#!/usr/bin/env python3
"""Deterministic TXT -> DOCX formatter for vocalized hadith text.

The transformer is intentionally conservative. It removes only proven footnote
blocks/page furniture, deletes parenthesized numeric footnote references, and
replaces the exact salutation phrase configured by the user. All other Arabic
letters and combining marks are preserved.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable, Sequence

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ARABIC_DIACRITICS_RE = re.compile(
    "[\u0610-\u061a\u064b-\u065f\u0670\u06d6-\u06ed]"
)
SEPARATOR_RE = re.compile(r"^\s*_{10,}\s*$")
PAGE_FURNITURE_RE = re.compile(
    r"^\s*.*(?:\(|\[)\s*ص\s*[:：]\s*\d+\s*(?:\)|\])\s*$"
)
FOOTNOTE_REF_RE = re.compile(r"\(\s*\d+\s*\)")
FOOTNOTE_START_RE = re.compile(r"^\s*\(\s*\d+\s*\)")
HADITH_START_RE = re.compile(r"^\s*(\d+)\s*-\s+")
RADI_RE = re.compile(r"رضي الله (?:عنهما|عنهم|عنهن|عنه|عنها)")
MARKER_RE = re.compile(r"HFU(\d{6})([SE])")
PART_MARKER_RE = re.compile(r"HFPART(\d{3})")

STOP_WORDS = {
    "عن", "قال", "قالت", "قالوا", "يقول", "سمعت", "حدثنا", "حدثني",
    "اخبرنا", "اخبرني", "ان", "انه", "انها", "رأى", "راى", "شهدت",
    "سأل", "سال", "مولى", "رسول", "النبي", "وفي", "وله", "ولمسلم",
    "في", "حديث",
}

PART_ORDINALS = {
    1: "الأَوَّلُ", 2: "الثَّانِي", 3: "الثَّالِثُ", 4: "الرَّابِعُ",
    5: "الخَامِسُ", 6: "السَّادِسُ", 7: "السَّابِعُ", 8: "الثَّامِنُ",
    9: "التَّاسِعُ", 10: "العَاشِرُ", 11: "الحَادِي عَشَرَ",
    12: "الثَّانِي عَشَرَ", 13: "الثَّالِثَ عَشَرَ", 14: "الرَّابِعَ عَشَرَ",
    15: "الخَامِسَ عَشَرَ", 16: "السَّادِسَ عَشَرَ", 17: "السَّابِعَ عَشَرَ",
    18: "الثَّامِنَ عَشَرَ", 19: "التَّاسِعَ عَشَرَ", 20: "العِشْرُونَ",
    21: "الحَادِي وَالعِشْرُونَ", 22: "الثَّانِي وَالعِشْرُونَ",
    23: "الثَّالِثُ وَالعِشْرُونَ", 24: "الرَّابِعُ وَالعِشْرُونَ",
    25: "الخَامِسُ وَالعِشْرُونَ", 26: "السَّادِسُ وَالعِشْرُونَ",
    27: "السَّابِعُ وَالعِشْرُونَ", 28: "الثَّامِنُ وَالعِشْرُونَ",
    29: "التَّاسِعُ وَالعِشْرُونَ", 30: "الثَّلَاثُونَ",
}


@dataclass
class Unit:
    index: int
    kind: str  # title | hadith | body
    text: str
    hadith_number: str | None = None


@dataclass
class CleanStats:
    page_markers_removed: int = 0
    separators_removed: int = 0
    footnote_blocks_removed: int = 0
    footnote_reference_numbers_removed: int = 0
    salutation_replacements: int = 0


@dataclass
class LayoutResult:
    page_count: int
    boundaries: list[int]
    padding_pages: list[int]
    part_start_pages: list[int]
    unit_pages: dict[int, tuple[int, int]]
    warnings: list[str]


def strip_diacritics(text: str) -> str:
    return ARABIC_DIACRITICS_RE.sub("", text)


def normalized_word(text: str) -> str:
    text = strip_diacritics(text)
    text = text.replace("ٱ", "ا").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    return re.sub(r"[^\u0621-\u064a]", "", text)


def read_utf8(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("The input TXT must be UTF-8 encoded.") from exc


def _process_page_chunk(lines: list[str], stats: CleanStats) -> list[str]:
    """Remove the last proven footnote block in a page chunk."""
    sep_positions = [i for i, line in enumerate(lines) if SEPARATOR_RE.match(line)]
    kept = lines
    if sep_positions:
        last_sep = sep_positions[-1]
        after = lines[last_sep + 1 :]
        first_nonblank = next((x for x in after if x.strip()), "")
        if first_nonblank and FOOTNOTE_START_RE.match(first_nonblank):
            stats.footnote_blocks_removed += 1
            kept = lines[:last_sep]
    # Separators are page furniture even when one occurs at the beginning of a
    # page after the previous page's footer marker.
    return [line for line in kept if not _count_separator(line, stats)]


def _count_separator(line: str, stats: CleanStats) -> bool:
    if SEPARATOR_RE.match(line):
        stats.separators_removed += 1
        return True
    return False


def clean_source(text: str, replacements: dict[str, str]) -> tuple[str, CleanStats]:
    """Clean page footnotes while preserving the body character stream."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    stats = CleanStats()
    pages: list[list[str]] = []
    joins: list[bool] = []
    current: list[str] = []

    for line in text.split("\n"):
        if PAGE_FURNITURE_RE.match(line) and "ص" in line:
            pages.append(_process_page_chunk(current, stats))
            joins.append(True)
            current = []
            stats.page_markers_removed += 1
        else:
            current.append(line)
    pages.append(_process_page_chunk(current, stats))

    # Merge pages. A physical page transition is not a semantic paragraph break.
    merged_lines: list[str] = []
    for page_idx, page_lines in enumerate(pages):
        while page_lines and not page_lines[0].strip():
            page_lines.pop(0)
        while page_lines and not page_lines[-1].strip():
            page_lines.pop()
        if not page_lines:
            continue
        if merged_lines and page_idx > 0:
            first = page_lines[0]
            prev = merged_lines[-1] if merged_lines else ""
            first_is_new_unit = bool(HADITH_START_RE.match(first)) or is_title(first, ["كتاب", "باب", "فصل"])
            if prev.strip() and not first_is_new_unit and not is_subentry_line(first):
                merged_lines[-1] = prev.rstrip() + " " + first.lstrip()
                page_lines = page_lines[1:]
        merged_lines.extend(page_lines)

    cleaned = "\n".join(merged_lines)
    cleaned, count = FOOTNOTE_REF_RE.subn("", cleaned)
    stats.footnote_reference_numbers_removed = count
    for old, new in replacements.items():
        cleaned, count = re.subn(re.escape(old), lambda _: new, cleaned)
        stats.salutation_replacements += count

    # Remove only blank-line inflation; do not normalize Arabic or punctuation.
    cleaned = re.sub(r"\n[ \t]+\n", "\n\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip() + "\n"
    return cleaned, stats


def is_title(line: str, prefixes: Sequence[str]) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    plain = strip_diacritics(stripped)
    plain = plain.replace("ٱ", "ا").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    return any(plain.startswith(prefix) for prefix in prefixes)


def is_subentry_line(line: str) -> bool:
    plain = strip_diacritics(line.strip())
    plain = plain.replace("ٱ", "ا").replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    prefixes = ("وفي", "ولمسلم", "وله", "وللبخاري", "وفي رواية", "وفي لفظ")
    return plain.startswith(prefixes)


def _join_unit_lines(lines: list[str]) -> str:
    out: list[str] = []
    pending_break = False
    for line in lines:
        if not line.strip():
            pending_break = True
            continue
        piece = line.strip()
        if not out:
            out.append(piece)
        elif pending_break or is_subentry_line(piece) or (":" in piece[:30] and not piece.startswith("«")):
            out.append("\n" + piece)
        else:
            out.append(" " + piece)
        pending_break = False
    return "".join(out).strip()


def parse_units(cleaned: str, title_prefixes: Sequence[str]) -> list[Unit]:
    units: list[Unit] = []
    current_kind: str | None = None
    current_lines: list[str] = []
    current_num: str | None = None

    def flush() -> None:
        nonlocal current_kind, current_lines, current_num
        if current_kind and current_lines:
            text = _join_unit_lines(current_lines)
            if text:
                units.append(Unit(len(units), current_kind, text, current_num))
        current_kind = None
        current_lines = []
        current_num = None

    for line in cleaned.splitlines():
        if is_title(line, title_prefixes):
            flush()
            units.append(Unit(len(units), "title", line.strip(), None))
            continue
        m = HADITH_START_RE.match(line)
        if m:
            flush()
            current_kind = "hadith"
            current_num = m.group(1)
            current_lines = [line]
            continue
        if current_kind == "hadith":
            current_lines.append(line)
        elif line.strip():
            if current_kind != "body":
                flush()
                current_kind = "body"
            current_lines.append(line)
        elif current_kind:
            current_lines.append(line)
    flush()
    for i, unit in enumerate(units):
        unit.index = i
    return units


def compact_chars(text: str) -> str:
    return re.sub(r"\s+", "", text)


def audit_unit_preservation(cleaned: str, units: Sequence[Unit]) -> None:
    reconstructed = "\n".join(unit.text for unit in units)
    if compact_chars(cleaned) != compact_chars(reconstructed):
        raise RuntimeError("Preservation audit failed: parsed output changed non-whitespace characters.")


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_paragraph_bidi(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        ppr.insert_element_before(
            bidi,
            "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
            "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
            "w:jc", "w:textDirection", "w:textAlignment",
            "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
            "w:rPr", "w:sectPr", "w:pPrChange",
        )


def set_run_style(run, *, font: str, size_pt: float, color: str, bold: bool = False, rtl: bool = True) -> None:
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(int(round(size_pt * 2))))
    rtl_el = rpr.find(qn("w:rtl"))
    if rtl and rtl_el is None:
        rpr.append(OxmlElement("w:rtl"))
    if not rtl and rtl_el is not None:
        rpr.remove(rtl_el)


def add_debug_marker(paragraph, marker: str) -> None:
    run = paragraph.add_run(marker)
    set_run_style(run, font="Arial", size_pt=1, color="FFFFFF", bold=False, rtl=False)


def _token_spans(text: str) -> list[tuple[int, int, str]]:
    return [(m.start(), m.end(), m.group()) for m in re.finditer(r"[\u0600-\u06FF]+|[^\u0600-\u06FF\s]+", text)]


def companion_spans(text: str) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    tokens = _token_spans(text)
    for radi in RADI_RE.finditer(text):
        before = [t for t in tokens if t[1] <= radi.start()]
        if not before:
            continue
        start = before[-1][0]
        scanned = 0
        for s, e, token in reversed(before):
            scanned += 1
            normalized = normalized_word(token)
            if token in {":", ".", "؛", ";", "-", "–", "—", "«", "»"}:
                start = e
                break
            if normalized in STOP_WORDS:
                start = e
                break
            start = s
            if scanned >= 24:
                break
        while start < radi.start() and text[start] in " \t\n،,:;-–—":
            start += 1
        if start < radi.start():
            spans.append((start, radi.start()))
    return _merge_spans(spans)


def matn_spans(text: str) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    stack: list[int] = []
    for i, char in enumerate(text):
        if char == "«":
            stack.append(i)
        elif char == "»" and stack:
            start = stack.pop()
            if not stack:
                spans.append((start, i + 1))
    return spans


def _merge_spans(spans: Sequence[tuple[int, int]]) -> list[tuple[int, int]]:
    if not spans:
        return []
    ordered = sorted(spans)
    merged = [list(ordered[0])]
    for start, end in ordered[1:]:
        if start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return [(a, b) for a, b in merged]


def style_segments(text: str) -> list[tuple[str, str, bool]]:
    """Return (text, role, bold) segments with deterministic precedence."""
    number_match = re.match(r"^\s*\d+\s*-", text)
    number = (number_match.start(), number_match.end()) if number_match else None
    companions = companion_spans(text)
    matn = matn_spans(text)
    bounds = {0, len(text)}
    if number:
        bounds.update(number)
    for s, e in companions + matn:
        bounds.add(s); bounds.add(e)
    ordered = sorted(bounds)
    result: list[tuple[str, str, bool]] = []
    for a, b in zip(ordered, ordered[1:]):
        if a == b:
            continue
        role = "isnad"
        bold = False
        if any(a >= s and b <= e for s, e in matn):
            role, bold = "matn", True
        elif number and a >= number[0] and b <= number[1]:
            role, bold = "number", True
        elif any(a >= s and b <= e for s, e in companions):
            role = "companion"
        chunk = text[a:b]
        if result and result[-1][1:] == (role, bold):
            result[-1] = (result[-1][0] + chunk, role, bold)
        else:
            result.append((chunk, role, bold))
    return result


def set_cell_margins(cell, top=60, start=0, bottom=60, end=0) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.insert_element_before(
            tc_mar,
            "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
            "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
        )
    for m, val in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.insert_element_before(
            borders,
            "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
            "w:tblCaption", "w:tblDescription", "w:tblPrChange",
        )
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "none")
        tag.set(qn("w:sz"), "0")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "FFFFFF")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph, config: dict) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bidi(paragraph)
    for kind, text in (("begin", None), (None, "PAGE"), ("separate", None), (None, "1"), ("end", None)):
        run = paragraph.add_run()
        set_run_style(run, font=config["font_name"], size_pt=config["footer_font_size_pt"], color="000000", rtl=True)
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        elif text == "PAGE":
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = " PAGE "
            run._r.append(instr)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def configure_section(section, config: dict, part_number: int, debug_markers: bool, first: bool) -> None:
    p = config["page"]
    section.page_width = Inches(p["width_inches"])
    section.page_height = Inches(p["height_inches"])
    section.top_margin = Inches(p["top_margin_inches"])
    section.bottom_margin = Inches(p["bottom_margin_inches"])
    section.left_margin = Inches(p["left_margin_inches"])
    section.right_margin = Inches(p["right_margin_inches"])
    section.header_distance = Inches(p["header_distance_inches"])
    section.footer_distance = Inches(p["footer_distance_inches"])

    header = section.header
    header.is_linked_to_previous = False
    # Clear all header content and reproduce the reference's blank line + title.
    for para in list(header.paragraphs):
        para._element.getparent().remove(para._element)
    blank = header.add_paragraph()
    blank.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bidi(blank)
    title_p = header.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bidi(title_p)
    run = title_p.add_run(f"الجُزْءُ {PART_ORDINALS.get(part_number, str(part_number))}")
    set_run_style(run, font=config["font_name"], size_pt=config["header_font_size_pt"], color=config["colors"]["part_header"], bold=True, rtl=True)
    if debug_markers:
        add_debug_marker(title_p, f"HFPART{part_number:03d}")

    footer = section.footer
    if not first:
        footer.is_linked_to_previous = True
    else:
        footer.is_linked_to_previous = False
        for para in list(footer.paragraphs):
            para._element.getparent().remove(para._element)
        fp = footer.add_paragraph()
        add_page_field(fp, config)


def apply_document_settings(doc: Document, config: dict) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.insert_element_before(
            update,
            "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat",
            "w:docVars", "w:rsids", "m:mathPr", "w:attachedSchema",
            "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats",
            "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
            "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary",
            "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator",
        )
    update.set(qn("w:val"), "true")
    zoom = settings.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    theme_lang = settings.find(qn("w:themeFontLang"))
    if theme_lang is None:
        theme_lang = OxmlElement("w:themeFontLang")
        settings.append(theme_lang)
    theme_lang.set(qn("w:bidi"), "ar-SA")


def add_title(doc: Document, unit: Unit, config: dict, debug_markers: bool) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_bidi(p)
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(15)
    p.paragraph_format.keep_with_next = True
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), config["colors"]["title_fill"])
    ppr.insert_element_before(
        shd,
        "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
        "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
        "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
        "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )
    if debug_markers:
        add_debug_marker(p, f"HFU{unit.index:06d}S")
    run = p.add_run(unit.text)
    set_run_style(run, font=config["font_name"], size_pt=config["title_font_size_pt"], color=config["colors"]["title_text"], bold=True, rtl=True)
    if debug_markers:
        add_debug_marker(p, f"HFU{unit.index:06d}E")


def new_hadith_table(doc: Document):
    table = doc.add_table(rows=0, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_no_borders(table)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), "9400")
    tbl_w.set(qn("w:type"), "dxa")
    return table


def add_text_unit(table, unit: Unit, config: dict, debug_markers: bool) -> None:
    row = table.add_row()
    set_row_cant_split(row)
    cell = row.cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)
    # Keep only one paragraph.
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_bidi(p)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.line_spacing = config["line_spacing"]
    p.paragraph_format.space_after = Pt(config["paragraph_space_after_pt"])
    if debug_markers:
        add_debug_marker(p, f"HFU{unit.index:06d}S")
    if unit.kind == "hadith":
        for chunk, role, bold in style_segments(unit.text):
            run = p.add_run(chunk)
            set_run_style(run, font=config["font_name"], size_pt=config["body_font_size_pt"], color=config["colors"][role], bold=bold, rtl=True)
    else:
        run = p.add_run(unit.text)
        set_run_style(run, font=config["font_name"], size_pt=config["body_font_size_pt"], color=config["colors"]["isnad"], bold=False, rtl=True)
    if debug_markers:
        add_debug_marker(p, f"HFU{unit.index:06d}E")


def build_document(
    units: Sequence[Unit],
    config: dict,
    output: Path,
    boundaries: Sequence[int] | None = None,
    padding_pages: Sequence[int] | None = None,
    debug_markers: bool = False,
) -> None:
    boundaries = list(boundaries or [])
    padding_pages = list(padding_pages or [0] * len(boundaries))
    if len(padding_pages) != len(boundaries):
        raise ValueError("padding_pages must match boundaries")
    boundary_map = {idx: pos for pos, idx in enumerate(boundaries)}

    doc = Document()
    # Remove the automatically-created empty paragraph from body.
    body = doc._element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)
    apply_document_settings(doc, config)
    configure_section(doc.sections[0], config, 1, debug_markers, first=True)

    table = None
    part_num = 1
    for unit in units:
        if unit.index in boundary_map:
            bpos = boundary_map[unit.index]
            table = None
            for _ in range(padding_pages[bpos]):
                p = doc.add_paragraph()
                p.add_run().add_break(WD_BREAK.PAGE)
            part_num += 1
            sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
            configure_section(sec, config, part_num, debug_markers, first=False)
        if unit.kind == "title":
            table = None
            add_title(doc, unit, config, debug_markers)
        else:
            if table is None:
                table = new_hadith_table(doc)
            add_text_unit(table, unit, config, debug_markers)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def find_soffice() -> str | None:
    candidates = ["libreoffice", "soffice"]
    for candidate in candidates:
        path = shutil.which(candidate)
        if path:
            return path
    return None


def render_pdf(docx_path: Path, workdir: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice/soffice is required for exact 20-page part pagination.")
    outdir = workdir / "pdf"
    outdir.mkdir(parents=True, exist_ok=True)
    profile = workdir / "lo_profile"
    home = workdir / "home"
    profile.mkdir(parents=True, exist_ok=True)
    home.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["HOME"] = str(home)
    cmd = [
        soffice, "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path),
    ]
    proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, env=env, timeout=180)
    pdf = outdir / f"{docx_path.stem}.pdf"
    if proc.returncode != 0 or not pdf.exists() or pdf.stat().st_size == 0:
        raise RuntimeError(f"LibreOffice PDF conversion failed: {proc.stdout}\n{proc.stderr}")
    return pdf


def extract_pdf_markers(pdf_path: Path) -> tuple[int, dict[int, tuple[int, int]], list[int]]:
    if fitz is None:
        raise RuntimeError("PyMuPDF is required for pagination analysis.")
    pdf = fitz.open(pdf_path)
    starts: dict[int, int] = {}
    ends: dict[int, int] = {}
    part_starts: dict[int, int] = {}
    for page_idx, page in enumerate(pdf, start=1):
        text = page.get_text("text")
        compact = re.sub(r"\s+", "", text)
        for m in MARKER_RE.finditer(compact):
            idx = int(m.group(1)); kind = m.group(2)
            if kind == "S":
                starts.setdefault(idx, page_idx)
            else:
                ends[idx] = page_idx
        for m in PART_MARKER_RE.finditer(compact):
            part_starts.setdefault(int(m.group(1)), page_idx)
    unit_pages = {idx: (starts[idx], ends.get(idx, starts[idx])) for idx in starts}
    ordered_part_starts = [part_starts[k] for k in sorted(part_starts)]
    return len(pdf), unit_pages, ordered_part_starts


def choose_initial_boundaries(unit_pages: dict[int, tuple[int, int]], page_count: int, part_pages: int, unit_count: int) -> list[int]:
    boundaries: list[int] = []
    for desired in range(part_pages + 1, page_count + 1, part_pages):
        candidate = None
        for idx in range(unit_count):
            if idx not in unit_pages:
                continue
            start, end = unit_pages[idx]
            if start < desired <= end:
                candidate = idx  # move a crossing oversized unit to the new part
                break
            if start >= desired:
                candidate = idx
                break
        if candidate is not None and candidate > 0 and candidate not in boundaries:
            boundaries.append(candidate)
    return boundaries


def resolve_layout(units: Sequence[Unit], config: dict, workdir: Path) -> LayoutResult:
    part_pages = int(config["part_pages"])
    provisional = workdir / "provisional.docx"
    build_document(units, config, provisional, debug_markers=True)
    pdf = render_pdf(provisional, workdir / "provisional_render")
    page_count, unit_pages, _ = extract_pdf_markers(pdf)
    if len(unit_pages) != len(units):
        missing = sorted(set(range(len(units))) - set(unit_pages))
        raise RuntimeError(f"Pagination markers missing for units: {missing}")
    boundaries = choose_initial_boundaries(unit_pages, page_count, part_pages, len(units))
    padding = [0] * len(boundaries)
    warnings: list[str] = []

    for iteration in range(40):
        debug_docx = workdir / f"layout-{iteration:02d}.docx"
        build_document(units, config, debug_docx, boundaries, padding, debug_markers=True)
        pdf = render_pdf(debug_docx, workdir / f"layout-render-{iteration:02d}")
        page_count, current_unit_pages, part_starts = extract_pdf_markers(pdf)
        required_parts = max(1, math.ceil(page_count / part_pages))

        if required_parts > len(boundaries) + 1:
            desired = (len(boundaries) + 1) * part_pages + 1
            candidate = None
            for idx in range(len(units)):
                if idx not in current_unit_pages:
                    continue
                start, end = current_unit_pages[idx]
                if start < desired <= end or start >= desired:
                    candidate = idx
                    break
            if candidate is not None and candidate > (boundaries[-1] if boundaries else 0):
                boundaries.append(candidate)
                padding.append(0)
                continue

        changed = False
        for k, boundary in enumerate(list(boundaries), start=1):
            desired = k * part_pages + 1
            actual = part_starts[k] if len(part_starts) > k else None
            if actual is None:
                raise RuntimeError(f"Missing section marker for part {k + 1}")
            if actual < desired:
                padding[k - 1] += desired - actual
                changed = True
                break
            if actual > desired:
                if padding[k - 1] > 0:
                    reduce_by = min(padding[k - 1], actual - desired)
                    padding[k - 1] -= reduce_by
                    changed = True
                    break
                min_idx = boundaries[k - 2] + 1 if k > 1 else 1
                if boundaries[k - 1] > min_idx:
                    boundaries[k - 1] -= 1
                    changed = True
                    break
                warnings.append(
                    f"Part {k + 1} starts on page {actual}, not {desired}; an indivisible unit prevents an earlier boundary."
                )
        if changed:
            continue
        return LayoutResult(page_count, boundaries, padding, part_starts, current_unit_pages, warnings)
    raise RuntimeError("Could not stabilize 20-page part boundaries after 40 iterations.")


def write_clean_text(units: Sequence[Unit], path: Path) -> None:
    blocks = [unit.text for unit in units]
    path.write_text("\n\n".join(blocks).strip() + "\n", encoding="utf-8")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Format vocalized hadith TXT as a reference-matched Word document.")
    parser.add_argument("input_txt", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--config", type=Path, default=Path(__file__).resolve().parents[1] / "config" / "default.json")
    parser.add_argument("--audit-json", type=Path)
    parser.add_argument("--clean-txt", type=Path)
    parser.add_argument("--workdir", type=Path)
    parser.add_argument("--skip-pagination", action="store_true", help="Create one-part output without exact 20-page section resolution.")
    args = parser.parse_args(argv)

    config = json.loads(args.config.read_text(encoding="utf-8"))
    original = read_utf8(args.input_txt)
    cleaned, stats = clean_source(original, config["replace_exact"])
    units = parse_units(cleaned, config["title_prefixes_unvocalized"])
    if not units:
        raise RuntimeError("No content units were detected.")
    audit_unit_preservation(cleaned, units)

    workdir_obj = None
    if args.workdir:
        args.workdir.mkdir(parents=True, exist_ok=True)
        workdir = args.workdir
    else:
        workdir_obj = tempfile.TemporaryDirectory(prefix="hadith_formatter_")
        workdir = Path(workdir_obj.name)

    if args.skip_pagination:
        layout = LayoutResult(0, [], [], [1], {}, ["Pagination was skipped by request."])
        build_document(units, config, args.output_docx, debug_markers=False)
    else:
        layout = resolve_layout(units, config, workdir)
        build_document(units, config, args.output_docx, layout.boundaries, layout.padding_pages, debug_markers=False)
        final_pdf = render_pdf(args.output_docx, workdir / "final_render")
        if fitz is not None:
            layout.page_count = len(fitz.open(final_pdf))

    clean_path = args.clean_txt or args.output_docx.with_suffix(".clean.txt")
    write_clean_text(units, clean_path)

    hadith_numbers = [u.hadith_number for u in units if u.kind == "hadith"]
    audit = {
        "skill": config["skill_name"],
        "version": config["version"],
        "input": str(args.input_txt),
        "output": str(args.output_docx),
        "clean_text": str(clean_path),
        "input_sha256": sha256_text(original),
        "cleaned_non_whitespace_sha256": sha256_text(compact_chars(cleaned)),
        "output_units_non_whitespace_sha256": sha256_text(compact_chars("\n".join(u.text for u in units))),
        "preservation_check": "PASS",
        "stats": asdict(stats),
        "units": {
            "total": len(units),
            "titles": sum(u.kind == "title" for u in units),
            "hadiths": sum(u.kind == "hadith" for u in units),
            "body_blocks": sum(u.kind == "body" for u in units),
            "hadith_numbers": hadith_numbers,
        },
        "layout": asdict(layout),
        "rules": {
            "font": config["font_name"],
            "font_size_pt": config["body_font_size_pt"],
            "part_pages": config["part_pages"],
            "salutation_replacement_only": config["replace_exact"],
            "companion_supplications_preserved": True,
        },
    }
    audit_path = args.audit_json or args.output_docx.with_suffix(".audit.json")
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"docx": str(args.output_docx), "clean_txt": str(clean_path), "audit": str(audit_path)}, ensure_ascii=False))
    if workdir_obj is not None:
        workdir_obj.cleanup()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
