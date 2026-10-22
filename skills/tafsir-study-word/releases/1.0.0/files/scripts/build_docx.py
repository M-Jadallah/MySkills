"""Build the editable green/gold tafsir Word document from reviewed study JSON.

This is a typesetter, not a tafsir extractor or vocalization engine.
Run audit_study.py before building and after updating Word fields.
"""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ASSETS = Path(__file__).resolve().parents[1] / 'assets'
AR = str.maketrans('0123456789', '٠١٢٣٤٥٦٧٨٩')
AXES = [
    'خُلَاصَةُ التَّفْسِيرِ',
    'جَمْعُ أَقْوَالِ الْمُفَسِّرِينَ وَتَرْجِيحَاتِهِمْ وَأَسْبَابِهَا',
    'الْمُصْطَلَحَاتُ الْعِلْمِيَّةُ وَشَرْحُهَا',
    'الْإِسْرَائِيلِيَّاتُ الْوَارِدَةُ فِي التَّفَاسِيرِ',
    'الْمَوْضُوعَاتُ الْمُسْتَخْلَصَةُ مِنَ الْآيَةِ وَالشَّوَاهِدُ عَلَيْهَا',
    'الْفَوَائِدُ التَّرْبَوِيَّةُ وَالْأَخْلَاقِيَّةُ الْمُسْتَنْبَطَةُ مِنَ الْآيَةِ',
    'تَفْسِيرُ أَسْمَاءِ اللَّهِ تَعَالَى وَصِفَاتِهِ الْوَارِدَةِ فِي الْآيَةِ',
]


def ar(value):
    return str(value).translate(AR)


def el(name, **attrs):
    node = OxmlElement(name)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def set_child(parent, name, **attrs):
    old = parent.find(qn(name))
    if old is not None:
        parent.remove(old)
    new = el(name, **attrs)
    parent.append(new)
    return new


class Builder:
    def __init__(self, study, spec, reference=None):
        self.study = study
        self.spec = spec
        self.colors = spec['colors']
        self.sources = {s['id']: s for s in study['sources']}
        # Preserve the selected, retained Word design rather than a host's
        # unrelated default.docx. Remove only its editable body content.
        self.reference = Path(reference) if reference else ASSETS / 'reference.docx'
        if not self.reference.is_file():
            raise FileNotFoundError('Retained Word design is missing: ' + str(self.reference))
        self.doc = Document(self.reference)
        for child in list(self.doc.element.body):
            if child.tag != qn('w:sectPr'):
                self.doc.element.body.remove(child)
        # Body hyperlinks belonged to the old study. Header/footer/image
        # relationships are preserved; new citations receive new links.
        for rid, rel in list(self.doc.part.rels.items()):
            if rel.reltype == RT.HYPERLINK:
                del self.doc.part.rels[rid]
        self.quotes = {q['id']: q for v in study['verses'] for g in v['groups'] for q in g['quotes']}
        self.setup()

    def rpr(self, parent, size, color, bold=False):
        fonts = set_child(parent, 'w:rFonts')
        for group in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + group), self.spec['font'])
        set_child(parent, 'w:sz', val=int(size * 2))
        set_child(parent, 'w:szCs', val=int(size * 2))
        set_child(parent, 'w:color', val=color)
        set_child(parent, 'w:rtl', val='1')
        set_child(parent, 'w:lang', val=self.spec['language'], bidi=self.spec['language'])
        set_child(parent, 'w:b', val='1' if bold else '0')
        set_child(parent, 'w:bCs', val='1' if bold else '0')

    def style(self, name, size, color, bold=False, line=29):
        if name not in self.doc.styles:
            self.doc.styles.add_style(name, 1)
        style = self.doc.styles[name]
        self.rpr(style.element.get_or_add_rPr(), size, color, bold)
        pf = style.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.line_spacing = Pt(line)
        pf.space_after = Pt(self.spec['body_space_after_pt'])
        pf.widow_control = True
        return style

    def setup(self):
        s, c = self.spec, self.colors
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = [Cm(x) for x in s['page_cm']]
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(s['margins_cm'])
        sec.header_distance = sec.footer_distance = Cm(s['header_footer_distance_cm'])
        set_child(sec._sectPr, 'w:bidi', val='1')
        for name in ('Normal', 'Body Text', 'Quotation', 'Title', 'Subtitle', 'Caption'):
            self.style(name, s['body_pt'], c['ink'], line=s['body_line_at_least_pt'])
        self.style('Title', s['title_pt'], c['green'], True).paragraph_format.space_after = Pt(18)
        # The bundled Word starter can carry an unrelated blue title rule.
        for name in ('Title', 'Title Char'):
            if name in self.doc.styles:
                for border in list(self.doc.styles[name].element.iter(qn('w:pBdr'))):
                    border.getparent().remove(border)
        self.style('Subtitle', s['subtitle_pt'], c['gold'], True)
        for level, size in enumerate(s['heading_pt'], 1):
            st = self.style('Heading ' + str(level), size, c['green'], True, s['heading_line_at_least_pt'])
            pf = st.paragraph_format
            pf.space_before, pf.space_after = Pt(14 if level == 2 else 9), Pt(6)
            pf.keep_with_next = pf.keep_together = True
            pf.page_break_before = level == 1
        self.style('Citation', s['citation_pt'], c['muted'], line=s['citation_line_at_least_pt']).paragraph_format.space_after = Pt(3)
        self.style('Term label', 22, c['green'], True).paragraph_format.keep_with_next = True
        verse = self.style('Verse', s['verse_pt'], c['green'], True, s['verse_line_at_least_pt'])
        verse.paragraph_format.space_before, verse.paragraph_format.space_after = Pt(10), Pt(20)
        quote = self.doc.styles['Quotation']
        quote.paragraph_format.right_indent = quote.paragraph_format.left_indent = Cm(s['quotation_indent_cm'])
        quote.paragraph_format.space_after = Pt(8)
        for name in ('TOC 1', 'TOC Heading'):
            toc = self.style(name, s['toc_pt'], c['green'], line=s['toc_line_at_least_pt'])
            pf = toc.paragraph_format
            pf.space_before = pf.space_after = Pt(0)
            pf.keep_with_next = pf.page_break_before = False
        set_child(self.doc.settings.element, 'w:updateFields', val='true')
        border = set_child(sec._sectPr, 'w:pgBorders', offsetFrom='page')
        if not 0 <= s['border_space_points'] <= 31:
            raise ValueError('Word page border space must be in 0..31 points')
        for side in ('top', 'left', 'bottom', 'right'):
            border.append(el('w:' + side, val='single', sz=s['border_eighth_points'], space=s['border_space_points'], color=c['gold']))
        # Keep the reference's repeating furniture and image relationships.
        if len(sec.header.tables) != 1 or len(sec.footer.tables) != 1:
            raise ValueError('Reference must have one three-cell header/footer row')
        hp = sec.header.tables[0].cell(0, 1).paragraphs[0]
        hp.clear()
        self.run(hp, self.study.get('header', self.study['title'] + '  •  دِرَاسَةٌ تَفْسِيرِيَّةٌ مُقَارَنَةٌ'), 16, c['green'])
        self.doc.core_properties.title = self.study['title']
        self.doc.core_properties.subject = self.study.get('subtitle', '')
        self.doc.core_properties.author = ''
        self.doc.core_properties.keywords = ''
        self.doc.core_properties.last_modified_by = ''
        self.doc.core_properties.language = self.spec['language']

    def run(self, paragraph, text, size=None, color=None, bold=None):
        run = paragraph.add_run(text)
        # Leave size/color/bold inherited unless explicitly overridden.
        rp = run._r.get_or_add_rPr()
        fonts = set_child(rp, 'w:rFonts')
        for group in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + group), self.spec['font'])
        set_child(rp, 'w:rtl', val='1')
        set_child(rp, 'w:lang', val=self.spec['language'], bidi=self.spec['language'])
        if size is not None:
            set_child(rp, 'w:sz', val=int(size * 2))
            set_child(rp, 'w:szCs', val=int(size * 2))
        if color:
            set_child(rp, 'w:color', val=color)
        if bold is not None:
            set_child(rp, 'w:b', val='1' if bold else '0')
            set_child(rp, 'w:bCs', val='1' if bold else '0')
        return run

    def text(self, text, style=None, center=False, color=None, size=None, bold=None):
        p = self.doc.add_paragraph(style=style)
        set_child(p._p.get_or_add_pPr(), 'w:bidi', val='1')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        self.run(p, text, size, color, bold)
        return p

    def citation(self, text, url=None):
        p = self.text(text, 'Citation')
        if url:
            run = p.runs[0]._r
            p._p.remove(run)
            h = el('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
            h.append(run)
            p._p.append(h)
        return p

    def line(self, p):
        border = set_child(p._p.get_or_add_pPr(), 'w:pBdr')
        border.append(el('w:bottom', val='single', sz='4', space='5', color=self.colors['gold']))

    def shade(self, p, border=False):
        pp = p._p.get_or_add_pPr()
        set_child(pp, 'w:shd', val='clear', fill=self.colors['pale'])
        if border:
            bd = set_child(pp, 'w:pBdr')
            for side in ('top', 'left', 'bottom', 'right'):
                bd.append(el('w:' + side, val='single', sz='5', space='9', color=self.colors['gold']))

    def heading(self, text, level=2):
        p = self.text(text, 'Heading ' + str(level), center=level == 1)
        if level == 2:
            self.shade(p)
            self.line(p)
        return p

    def field(self, p, instruction, display=''):
        run = self.run(p, '')
        run._r.append(el('w:fldChar', fldCharType='begin'))
        instr = el('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = ' ' + instruction + ' '
        run._r.append(instr)
        run._r.append(el('w:fldChar', fldCharType='separate'))
        if display:
            self.run(p, display)
        p.add_run()._r.append(el('w:fldChar', fldCharType='end'))

    def furniture(self, part, sy):
        widths = self.spec['header_table_columns_cm']
        table = part.add_table(rows=1, cols=3, width=Cm(sum(widths)))
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table._tbl.tblPr.append(el('w:bidiVisual', val='0'))
        borders = el('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(el('w:' + side, val='nil'))
        table._tbl.tblPr.append(borders)
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for i, cell in enumerate(table.rows[0].cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = el('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                margins.append(el('w:' + side, w='0', type='dxa'))
            cell._tc.get_or_add_tcPr().append(margins)
            p = cell.paragraphs[0]
            set_child(p._p.get_or_add_pPr(), 'w:bidi', val='1')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.AT_LEAST, Pt(30)
            if i == 1:
                if sy == 1:
                    self.run(p, self.study.get('header', self.study['title'] + '  •  دِرَاسَةٌ تَفْسِيرِيَّةٌ مُقَارَنَةٌ'), 16, self.colors['green'])
                else:
                    self.field(p, 'PAGE')
                    self.rpr(p.runs[0]._r.get_or_add_rPr(), 16, self.colors['green'])
            else:
                sx = 1 if i == 0 else -1
                p.add_run().add_picture(str(ASSETS / f'corner-{sx}-{sy}.png'), width=Cm(1), height=Cm(1))
        for p in part.paragraphs:
            pf = p.paragraph_format
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing = Pt(1)
            pf.keep_with_next = False

    def quote(self, q, verse):
        source = self.sources[q['source_id']]
        caption = q.get('citation') or (source['name'] + '؛ ' + verse['heading'])
        if q.get('owner'):
            caption += '؛ الْقَائِلُ: ' + q['owner']
        self.citation(caption, q.get('url')).paragraph_format.keep_with_next = True
        self.text('«' + q['text'] + '»', 'Quotation')

    def empty(self, verse, axis):
        message = verse.get('empty_axes', {}).get(str(axis))
        if not message:
            raise ValueError(f'{verse["key"]}: missing axis {axis} and no reviewed empty_axes explanation')
        self.text(message)

    def content(self):
        d, c = self.study, self.colors
        self.text(d['title'], 'Title', center=True)
        subtitle = self.text(d.get('subtitle', 'دِرَاسَةٌ تَفْسِيرِيَّةٌ مُقَارَنَةٌ'), 'Subtitle', center=True)
        self.line(subtitle)
        p = self.text('اسْتِخْلَاصُ الْمَعَانِي وَتَحْرِيرُ الْأَقْوَالِ\nوَتَوْثِيقُ الشَّوَاهِدِ وَالْفَوَائِدِ', center=True, color=c['green'], size=25)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(55), Pt(28)
        counts = 'عَدَدُ الْآيَاتِ: ' + ar(len(d['verses'])) + '\nسَبْعَةُ مَحَاوِرَ لِكُلِّ آيَةٍ\nعَدَدُ الْمَلَفَّاتِ الْمَصْدَرِيَّةِ: ' + ar(len(d['sources']))
        self.text(counts, center=True, size=24)
        self.text('إِعْدَادٌ مِنَ النُّصُوصِ الْمُعْتَمَدَةِ', center=True, color=c['muted'], size=19).paragraph_format.space_before = Pt(45)
        self.text('النَّمُوذَجُ التُّرَاثِيُّ الْأَخْضَرُ', center=True, color=c['gold'], size=20)
        p = self.text('فِهْرِسُ الْآيَاتِ وَالْمَصَادِرِ', center=True, color=c['green'], size=28, bold=True)
        p.paragraph_format.page_break_before = True
        p = self.text('')
        self.field(p, 'TOC \\o "1-1" \\h \\z \\u', 'فِهْرِسُ الْمُحْتَوَيَاتِ')
        self.heading('مَنْهَجُ الدِّرَاسَةِ وَضَبْطُ النُّقُولِ', 1)
        for paragraph in d['methodology']:
            self.text(paragraph)
        for v in d['verses']:
            self.heading(v['heading'], 1)
            p = self.text('﴿' + v['text'] + '﴾  ' + ar(v['key'].split(':')[-1]), 'Verse', center=True)
            self.shade(p, border=True)
            self.heading('١  ' + AXES[0])
            summary = v['summary']
            self.text(summary['text']).paragraph_format.keep_with_next = True
            self.citation(summary.get('citation') or self.sources[summary['source_id']]['name'], summary.get('url'))
            self.heading('٢  ' + AXES[1])
            if not v['groups']:
                self.empty(v, 2)
            for group in v['groups']:
                self.heading(group['title'], 3)
                for q in group['quotes']:
                    self.quote(q, v)
                if group.get('note'):
                    self.text('تَحْرِيرُ الْخِلَافِ: ' + group['note'])
                for pref in group.get('preferences', []):
                    source = self.sources[pref['source_id']]
                    self.citation('الْمُرَجِّحُ: ' + source['author'] + '؛ الْمَصْدَرُ: ' + source['name']).paragraph_format.keep_with_next = True
                    self.text(pref['text'])
                    if pref.get('reason'):
                        self.text('سَبَبُ التَّرْجِيحِ: ' + pref['reason'])
            self.heading('٣  ' + AXES[2])
            if not v['terms']:
                self.empty(v, 3)
            for term in v['terms']:
                self.text(term['name'] + '  —  ' + term['field'], 'Term label')
                for sid in dict.fromkeys(self.quotes[qid]['source_id'] for qid in term['source_quote_ids']):
                    self.citation('مَوْضِعُ الْمُصْطَلَحِ: ' + self.sources[sid]['name'] + '؛ ' + v['heading']).paragraph_format.keep_with_next = True
                self.text('التَّعْرِيفُ: ' + term['definition'])
                self.text('الشَّرْحُ: ' + term['explanation'])
                self.text('مِثَالٌ خَارِجِيٌّ: ' + term['example']).paragraph_format.space_after = Pt(10)
            self.heading('٤  ' + AXES[3])
            if not v['stories']:
                self.empty(v, 4)
            for story in v['stories']:
                self.quote(story['quote'], v)
                self.text('تَصْنِيفُ الْخَبَرِ: ' + story['classification']).paragraph_format.keep_with_next = True
                self.text('حُكْمُ النَّقْلِ: ' + story['judgment'])
                if story.get('evidence'):
                    self.text(story['evidence'])
            self.heading('٥  ' + AXES[4])
            if not v['themes']:
                self.empty(v, 5)
            for theme in v['themes']:
                self.heading(theme['title'], 3)
                self.text(theme['explanation'])
                self.citation('شَوَاهِدُ الْقُرْآنِ').paragraph_format.keep_with_next = True
                for proof in theme['quran']:
                    self.text('﴿' + proof['text'] + '﴾', color=c['green']).paragraph_format.keep_with_next = True
                    self.citation(proof['citation'], proof.get('url'))
                self.citation('شَوَاهِدُ السُّنَّةِ').paragraph_format.keep_with_next = True
                for proof in theme['hadith']:
                    self.text('رَوَاهُ ' + proof['narrator'] + ': «' + proof['text'] + '».').paragraph_format.keep_with_next = True
                    label = ('مِنْ حَدِيثٍ فِي ' if proof.get('excerpt') else '') + proof['reference'] + '؛ ' + proof['grade']
                    if proof.get('grade_source'):
                        label += '؛ ' + proof['grade_source']
                    self.citation(label, proof.get('url'))
                self.citation('شَاهِدٌ شِعْرِيٌّ لِلِاسْتِئْنَاسِ بِالْمَعْنَى').paragraph_format.keep_with_next = True
                for proof in theme['poetry']:
                    p = self.text(proof['text'], center=True)
                    p.paragraph_format.keep_together = p.paragraph_format.keep_with_next = True
                    self.citation('؛ '.join(x for x in (proof.get('author'), proof.get('reference'), proof.get('attribution_note')) if x), proof.get('url'))
            self.heading('٦  ' + AXES[5])
            if not v['benefits']:
                self.empty(v, 6)
            for i, benefit in enumerate(v['benefits'], 1):
                label = 'مَنْصُوصَةٌ: ' if benefit['type'] == 'direct' else 'مُسْتَنْبَطَةٌ: '
                self.text(ar(i) + '  ' + label + benefit['text'])
            self.heading('٧  ' + AXES[6])
            name_paragraphs = v['names']['paragraphs']
            for i, paragraph in enumerate(name_paragraphs):
                self.text(paragraph).paragraph_format.keep_with_next = bool(v['names']['references']) or i < len(name_paragraphs) - 1
            for ref in v['names']['references']:
                self.citation(ref['text'], ref.get('url'))
        self.heading('الْمَصَادِرُ وَمَفَاتِيحُ الْعَزْوِ', 1)
        for i, source in enumerate(d['sources'], 1):
            self.text(ar(i) + '  ' + source['name']).paragraph_format.keep_with_next = True
            self.citation('الْمُؤَلِّفُ: ' + source['author'])
            if source.get('bibliographic_note'):
                self.citation(source['bibliographic_note'], source.get('url'))
        self.heading('الْمَرَاجِعُ الْمُكَمِّلَةُ وَمَصَادِرُ الشَّوَاهِدِ', 2)
        for ref in d['bibliography']:
            if ref.get('url'):
                self.citation(ref['text'], ref['url'])
            else:
                self.text(ref['text'])
        return self.doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--study', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--style', type=Path, default=ASSETS / 'style-spec.json')
    parser.add_argument('--reference', type=Path, default=ASSETS / 'reference.docx')
    args = parser.parse_args()
    if args.out.resolve() in {(ASSETS / 'reference.docx').resolve(), args.reference.resolve()}:
        parser.error('The retained reference is read-only; select a task output path')
    study = json.loads(args.study.read_text(encoding='utf-8-sig'))
    spec = json.loads(args.style.read_text(encoding='utf-8-sig'))
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc = Builder(study, spec, args.reference).content()
    doc.save(args.out)
    print(json.dumps({'path': str(args.out.resolve()), 'verses': len(study['verses']), 'paragraphs': len(doc.paragraphs), 'fields_updated': False}, ensure_ascii=False))


if __name__ == '__main__':
    main()
