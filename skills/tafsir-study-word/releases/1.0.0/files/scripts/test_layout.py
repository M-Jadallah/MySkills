"""Regression tests for source Quran mapping and the editable Word typesetter."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from docx.oxml.ns import qn
from build_docx import Builder, ASSETS
from prepare_quran import prepare
from test_tools import fixture


class LayoutAndQuranTests(unittest.TestCase):
    def test_retained_quran_map_has_verified_boundaries(self):
        meta = json.loads((ASSETS / 'quran-source.json').read_text(encoding='utf-8'))
        q = prepare(ASSETS / 'quran-simple-1.1.txt', meta)
        self.assertEqual(len(q), 6236)
        self.assertEqual(q['91:1'], 'وَالشَّمْسِ وَضُحَاهَا')
        self.assertEqual(q['91:15'], 'وَلَا يَخَافُ عُقْبَاهَا')
        self.assertTrue(q['1:1'].startswith('بِسْمِ'))
        self.assertFalse(q['9:1'].startswith('بِسْمِ'))
        self.assertFalse(q['2:1'].startswith('بِسْمِ'))

    def test_changed_quran_source_is_rejected(self):
        meta = json.loads((ASSETS / 'quran-source.json').read_text(encoding='utf-8'))
        with tempfile.TemporaryDirectory() as folder:
            changed = Path(folder) / 'quran.txt'
            changed.write_bytes((ASSETS / 'quran-simple-1.1.txt').read_bytes() + b'\nchanged')
            with self.assertRaisesRegex(ValueError, 'differs'):
                prepare(changed, meta)

    def test_builder_keeps_content_editable_and_semantic(self):
        with tempfile.TemporaryDirectory() as folder:
            study, _, _, _ = fixture(Path(folder))
            spec = json.loads((ASSETS / 'style-spec.json').read_text(encoding='utf-8'))
            doc = Builder(study, spec).content()
            paragraphs = doc.paragraphs
            self.assertTrue(any(p.text == study['verses'][0]['summary']['text'] for p in paragraphs))
            self.assertEqual(len([p for p in paragraphs if p.style.name == 'Heading 2' and p.text[:1] in '١٢٣٤٥٦٧']), 7)
            self.assertFalse(doc.tables, 'Body content must remain native paragraphs')
            self.assertEqual(len(doc.sections[0].header.tables), 1)
            self.assertEqual(len(doc.sections[0].footer.tables), 1)
            self.assertTrue(all(p._p.get_or_add_pPr().find(qn('w:bidi')) is not None for p in paragraphs))
            field_text = ' '.join(n.text or '' for n in doc.element.iter(qn('w:instrText')))
            self.assertIn('TOC', field_text)
            self.assertEqual(doc.styles['Normal'].element.get_or_add_rPr().find(qn('w:szCs')).get(qn('w:val')), '44')
            self.assertFalse(list(doc.styles['Title'].element.iter(qn('w:pBdr'))), 'Remove the starter template blue title rule')

    def test_invalid_word_border_is_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            study, _, _, _ = fixture(Path(folder))
            spec = json.loads((ASSETS / 'style-spec.json').read_text(encoding='utf-8'))
            spec['border_space_points'] = 36
            with self.assertRaisesRegex(ValueError, '0..31'):
                Builder(study, spec)

    def test_new_study_does_not_inherit_old_subject_keywords_or_links(self):
        with tempfile.TemporaryDirectory() as folder:
            study, _, _, _ = fixture(Path(folder))
            spec = json.loads((ASSETS / 'style-spec.json').read_text(encoding='utf-8'))
            doc = Builder(study, spec).content()
            self.assertEqual(doc.core_properties.keywords, '')
            self.assertEqual(doc.core_properties.title, study['title'])
            self.assertEqual(doc.core_properties.subject, study['subtitle'])
            links = [rel.target_ref for rel in doc.part.rels.values() if rel.is_external]
            self.assertNotIn('https://sunnah.com/bukhari:1043', links)


if __name__ == '__main__':
    unittest.main(verbosity=2)
