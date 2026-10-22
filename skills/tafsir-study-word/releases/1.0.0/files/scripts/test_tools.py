#!/usr/bin/env python3
"""Meaningful positive/negative tests for ingestion and source auditing."""
from __future__ import annotations

import copy
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from audit_study import Auditor, normalized
from ingest_sources import ingest, split_blocks


def fixture(folder: Path):
    first, second = folder / "المختصر.txt", folder / "المصدر الثاني.txt"
    first.write_bytes("الآية ١\r\nأشرق نورها.\r\n".encode("utf-8"))
    second.write_bytes("الآية ۱\r\nأظهر ضوءها.\r\n".encode("utf-8"))
    sources = ingest([first, second], surah=91)
    sources["sources"][0]["name"] = "الْمُخْتَصَرُ فِي تَفْسِيرِ الْقُرْآنِ الْكَرِيمِ"
    sources["sources"][1]["name"] = "تَفْسِيرٌ ثَانٍ"
    for source in sources["sources"]:
        source["author"] = "مُؤَلِّفٌ"

    def quoted(index, text, qid=None):
        source = sources["sources"][index]
        block = source["blocks"][0]
        start = block["header_end"]
        result = {"text": text, "source_id": source["id"], "block_id": block["id"],
                  "source_spans": [{"start": start, "end": len(source["raw_text"])}]}
        if qid:
            result["id"] = qid
        return result

    summary = quoted(0, "أَشْرَقَ نُورُهَا.")
    summary["mode"] = "literal"
    quote = quoted(1, "أَظْهَرَ ضَوْءَهَا.", "q1")
    verse = {"key": "91:1", "heading": "الْآيَةُ الْأُولَى", "text": "وَالشَّمْسِ وَضُحَاهَا",
             "summary": summary,
             "groups": [{"id": "g1", "title": "الْقَوْلُ", "quotes": [quote], "preferences": []}],
             "terms": [], "stories": [],
             "themes": [{"id": "t1", "title": "آيَاتُ اللَّهِ", "explanation": "يُسْتَدَلُّ بِهَا.",
                         "quran": [{"key": "91:1", "text": "وَالشَّمْسِ وَضُحَاهَا", "citation": "الشَّمْسُ: ١"}],
                         "hadith": [{"text": "إِنَّمَا الْأَعْمَالُ بِالنِّيَّاتِ.", "narrator": "عُمَرُ",
                                     "reference": "صَحِيحُ الْبُخَارِيِّ، ١", "grade": "صَحِيحٌ",
                                     "grade_source": "صَحِيحُ الْبُخَارِيِّ", "url": "https://sunnah.com/bukhari:1",
                                     "excerpt": True}],
                         "poetry": [{"text": "وَفِي كُلِّ شَيْءٍ لَهُ آيَةٌ تَدُلُّ عَلَى أَنَّهُ وَاحِدُ.",
                                     "author": "أَبُو الْعَتَاهِيَةِ", "reference": "بُلُوغُ الْأَرَبِ",
                                     "url": "https://usul.ai/ar/t/bulugh-arab-1/26",
                                     "attribution_note": "تُنْسَبُ أَيْضًا إِلَى الشَّافِعِيِّ."}]}],
             "benefits": [{"id": "f1", "text": "فَائِدَةٌ.", "type": "inferred", "evidence_ids": ["q1"]}],
             "names": {"paragraphs": ["لَا يَرِدُ اسْمٌ صَرِيحٌ فِي هَذِهِ الْآيَةِ."], "references": []},
             "empty_axes": {"3": "لَا يَرِدُ مُصْطَلَحٌ.", "4": "لَا تَرِدُ رِوَايَةٌ قَصَصِيَّةٌ."}}
    study = {"schema_version": 1, "title": "دِرَاسَةٌ", "subtitle": "تَفْسِيرِيَّةٌ",
             "surah": {"number": 91, "name": "الشَّمْسُ"}, "methodology": ["مَنْهَجٌ."],
             "sources": [{"id": s["id"], "name": s["name"], "author": s["author"]} for s in sources["sources"]],
             "verses": [verse], "bibliography": []}
    coverage = {"items": [{"id": "c1", "source_id": "s01", "block_id": "b001", "verse_keys": ["91:1"],
                           "kind": "summary", "status": "included", "output_ids": ["91:1"]},
                          {"id": "c2", "source_id": "s02", "block_id": "b001", "verse_keys": ["91:1"],
                           "kind": "opinion", "status": "included", "output_ids": ["q1"]}]}
    quran = {"91:1": verse["text"]}
    return study, sources, quran, coverage


class IngestionTests(unittest.TestCase):
    def test_original_offsets_crlf_and_both_digit_sets(self):
        raw = "مُقَدِّمَةٌ 😀\r\nالآيات ١–٢\r\nأول نص.\r\nالآية ۳\r\nآخر نص.\r\n"
        blocks = split_blocks(raw, 91)
        self.assertEqual([b["verses"] for b in blocks], [[], [1, 2], [3]])
        self.assertTrue(blocks[0]["requires_mapping"])
        self.assertEqual("".join(b["text"] for b in blocks), raw)
        for block in blocks:
            self.assertEqual(raw[block["start"]:block["end"]], block["text"])
        self.assertEqual(blocks[1]["verse_keys"], ["91:1", "91:2"])

    def test_unknown_heading_is_never_guessed(self):
        raw = "الآية الأولى\nنص\nالآية 2\nآخر\n"
        blocks = split_blocks(raw)
        self.assertTrue(blocks[0]["requires_mapping"])
        self.assertEqual(blocks[0]["verses"], [])
        self.assertEqual(blocks[1]["verses"], [2])
        self.assertEqual(blocks[1]["verse_keys"], [])

    def test_reversed_range_is_unmapped(self):
        self.assertTrue(split_blocks("الآيات ٣–١\nنص")[0]["requires_mapping"])

    def test_source_range_suffix_is_explicit_not_guessed(self):
        self.assertEqual(split_blocks("الآيات ۱–۸ (نطاق المصدر)\nنص")[0]["verses"], list(range(1, 9)))

    def test_repeated_body_is_flagged_even_when_verse_headings_differ(self):
        with tempfile.TemporaryDirectory() as tmp:
            file = Path(tmp) / "م.txt"
            file.write_text("الآية 1\nنَصٌّ مُكَرَّرٌ.\nالآية 2\nنَصٌّ مُكَرَّرٌ.\n", encoding="utf-8")
            result = ingest([file])
            self.assertEqual(len(result["exact_duplicate_blocks"]), 0)
            self.assertEqual(result["exact_duplicate_content"], [[{"source_id": "s01", "block_id": "b001"},
                                                                 {"source_id": "s01", "block_id": "b002"}]])

    def test_exact_duplicates_preserve_separate_sources(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            data = "الآية 1\r\nنص متكرر.\r\n".encode("utf-8")
            for name in ("مؤلف أ.txt", "مؤلف ب.txt"):
                (folder / name).write_bytes(data)
            result = ingest([folder / "مؤلف أ.txt", folder / "مؤلف ب.txt"])
            self.assertEqual(len(result["sources"]), 2)
            self.assertEqual(result["exact_duplicate_files"], [["s01", "s02"]])
            self.assertEqual(result["sources"][0]["sha256"], hashlib.sha256(data).hexdigest())
            self.assertIsNone(result["sources"][0]["author"])

    def test_utf8_bom_and_newlines_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            file = Path(tmp) / "م.txt"
            data = b"\xef\xbb\xbf" + "الآية ١\r\nنَصٌّ\r\n".encode("utf-8")
            file.write_bytes(data)
            source = ingest([file])["sources"][0]
            self.assertEqual(source["encoding"], "utf-8-sig")
            self.assertEqual(source["sha256"], hashlib.sha256(data).hexdigest())
            self.assertTrue(source["raw_text"].endswith("\r\n"))


class AuditTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.study, self.sources, self.quran, self.coverage = fixture(self.folder)

    def tearDown(self):
        self.temp.cleanup()

    def audit(self, docx=None):
        return Auditor(self.study, self.sources, self.quran, self.coverage).run(docx)

    def codes(self, docx=None):
        return {e["code"] for e in self.audit(docx)["errors"]}

    def test_valid_minimal_study(self):
        result = self.audit()
        self.assertTrue(result["passed"], result)

    def test_hamza_change_is_not_hidden_by_normalization(self):
        self.study["verses"][0]["groups"][0]["quotes"][0]["text"] = "اظهر ضوءها."
        self.assertIn("literal_mismatch", self.codes())
        self.assertNotEqual(normalized("أ"), normalized("ا"))
        self.assertNotEqual(normalized("ى"), normalized("ي"))

    def test_wrong_attribution_does_not_search_every_source(self):
        quote = self.study["verses"][0]["groups"][0]["quotes"][0]
        quote["source_id"] = "s01"
        self.assertIn("literal_mismatch", self.codes())

    def test_span_outside_named_block_is_rejected(self):
        quote = self.study["verses"][0]["groups"][0]["quotes"][0]
        quote["source_spans"][0]["end"] += 1
        self.assertIn("quote_span", self.codes())

    def test_hadith_without_grade_fails(self):
        del self.study["verses"][0]["themes"][0]["hadith"][0]["grade"]
        self.assertIn("required_text", self.codes())

    def test_missing_quran_genre_fails(self):
        self.study["verses"][0]["themes"][0]["quran"] = []
        self.assertIn("theme_missing_quran", self.codes())

    def test_missing_hadith_genre_fails(self):
        self.study["verses"][0]["themes"][0]["hadith"] = []
        self.assertIn("theme_missing_hadith", self.codes())

    def test_missing_poetry_genre_fails(self):
        self.study["verses"][0]["themes"][0]["poetry"] = []
        self.assertIn("theme_missing_poetry", self.codes())

    def set_primary_gap_and_fallback(self):
        block = self.sources["sources"][0]["blocks"][0]
        block["verse_keys"], block["verses"] = ["91:2"], [2]
        self.coverage["items"][0].update(verse_keys=["91:2"], status="not_relevant",
                                        justification="خَارِجُ الْآيَةِ الْمَدْرُوسَةِ.")
        source = self.sources["sources"][1]
        source["name"] = "تَفْسِيرُ السَّعْدِيِّ"
        self.study["sources"][1]["name"] = source["name"]
        summary = copy.deepcopy(self.study["verses"][0]["groups"][0]["quotes"][0])
        summary.pop("id")
        summary.update(mode="literal", fallback_reason="نَقْصٌ فِي تَغْطِيَةِ الْمِلَفِّ الْأَصْلِيِّ.")
        self.study["verses"][0]["summary"] = summary

    def test_existing_primary_file_gap_does_not_automatically_allow_fallback(self):
        self.set_primary_gap_and_fallback()
        self.assertIn("primary_coverage_gap", self.codes())

    def test_explicit_recorded_user_authorization_allows_gap_fallback(self):
        self.set_primary_gap_and_fallback()
        self.study["verses"][0]["summary"]["fallback_authorization"] = {
            "authority": "user", "instruction": "استعمل السعدي للآية الناقصة فقط.", "verse_keys": ["91:1"]}
        result = self.audit()
        self.assertTrue(result["passed"], result)

    def test_missing_axis_fails(self):
        del self.study["verses"][0]["names"]
        self.assertIn("missing_axis_7", self.codes())

    def test_unexplained_empty_term_axis_fails(self):
        del self.study["verses"][0]["empty_axes"]["3"]
        self.assertIn("missing_axis_3", self.codes())

    def test_unknown_range_must_be_resolved(self):
        block = self.sources["sources"][1]["blocks"][0]
        block["verse_keys"], block["verses"], block["requires_mapping"] = [], [], True
        self.assertIn("requires_mapping", self.codes())

    def test_empty_quran_box_fails(self):
        self.study["verses"][0]["text"] = ""
        self.assertIn("quran_missing", self.codes())

    def test_quran_diacritic_change_fails(self):
        self.study["verses"][0]["text"] = "والشمس وضحاها"
        self.assertIn("quran_mismatch", self.codes())

    def test_missing_source_coverage_fails(self):
        self.coverage["items"].pop()
        self.assertIn("uncovered_block_verse", self.codes())
        self.assertIn("quote_without_coverage", self.codes())

    def test_duplicate_cycle_fails(self):
        for index, item in enumerate(self.coverage["items"]):
            item.update(status="duplicate", justification="تكرار.", duplicate_of=f"c{2-index}")
        self.assertIn("duplicate_cycle", self.codes())

    def test_broken_term_evidence_fails(self):
        self.study["verses"][0]["terms"] = [{"id": "term1", "name": "مَجَازٌ", "field": "بَلَاغَةٌ",
                                            "definition": "تَعْرِيفٌ.", "explanation": "شَرْحٌ.",
                                            "example": "مِثَالٌ.", "source_quote_ids": ["missing"]}]
        self.assertIn("broken_evidence_link", self.codes())

    def test_changed_original_file_fails(self):
        Path(self.sources["sources"][0]["path"]).write_text("تَغْيِيرٌ", encoding="utf-8")
        self.assertIn("source_file_changed", self.codes())

    def make_docx(self, axes=7):
        from build_docx import Builder, ASSETS
        doc = Builder(self.study, json.loads((ASSETS / 'style-spec.json').read_text(encoding='utf-8'))).content()
        axis_paragraphs = [p for p in doc.paragraphs if p.style.name == 'Heading 2' and p.text[:1] in '١٢٣٤٥٦٧']
        for p in (axis_paragraphs[axes:] if axes < 7 else []):
            p._p.getparent().remove(p._p)
        file = self.folder / "study.docx"
        doc.save(file)
        return file

    def test_docx_has_all_text_and_seven_axes(self):
        self.assertTrue(self.audit(self.make_docx())["passed"])

    def test_docx_with_dropped_axis_fails(self):
        self.assertIn("docx_axis_count", self.codes(self.make_docx(axes=6)))

    def replace_docx_text(self, file, old, new, count=-1):
        with zipfile.ZipFile(file) as archive:
            parts = {n: archive.read(n) for n in archive.namelist()}
        xml = parts['word/document.xml'].decode('utf-8')
        parts['word/document.xml'] = xml.replace(old, new, count).encode('utf-8')
        with zipfile.ZipFile(file, "w") as archive:
            for name, data in parts.items():
                archive.writestr(name, data)

    def test_docx_quran_diacritic_loss_is_rejected(self):
        file = self.make_docx()
        self.replace_docx_text(file, "وَالشَّمْسِ وَضُحَاهَا", "والشمس وضحاها", count=1)
        self.assertIn("docx_quran_mismatch", self.codes(file))

    def test_docx_axis_number_order_is_checked(self):
        file = self.make_docx()
        self.replace_docx_text(file, "٣  الْمُصْطَلَحَاتُ", "٢  الْمُصْطَلَحَاتُ")
        self.assertIn("docx_axis_order", self.codes(file))

    def test_docx_names_paragraph_cannot_disappear(self):
        file = self.make_docx()
        self.replace_docx_text(file, self.study["verses"][0]["names"]["paragraphs"][0], "")
        self.assertIn("docx_missing_text", self.codes(file))

    def test_docx_non_quran_vowels_cannot_disappear(self):
        file = self.make_docx()
        self.replace_docx_text(file, 'أَظْهَرَ ضَوْءَهَا.', 'أظهر ضوءها.')
        self.assertIn('docx_missing_text', self.codes(file))

    def test_docx_hadith_grade_cannot_disappear(self):
        file = self.make_docx()
        self.replace_docx_text(file, '؛ صَحِيحٌ', '؛ ')
        self.assertTrue(any(e['code'] == 'docx_missing_text' and e['path'].endswith('.grade') for e in self.audit(file)['errors']))

    def test_docx_story_classification_is_printed_and_checked(self):
        v = self.study['verses'][0]
        q = copy.deepcopy(v['groups'][0]['quotes'][0]); q['id'] = 'storyquote'
        v['stories'] = [{'id': 'story1', 'quote': q, 'classification': 'خَبَرٌ قَصَصِيٌّ غَيْرُ مُحَقَّقٍ.', 'judgment': 'لَا يَثْبُتُ.', 'evidence': 'لَا يَكْفِي هَذَا الْمِلَفُّ لِتَصْحِيحِهِ.'}]
        self.coverage['items'][1]['output_ids'].append('storyquote')
        v['empty_axes'].pop('4')
        file = self.make_docx()
        self.assertTrue(self.audit(file)['passed'])
        self.replace_docx_text(file, v['stories'][0]['classification'], '')
        self.assertTrue(any(e['code'] == 'docx_missing_text' and e['path'].endswith('.classification') for e in self.audit(file)['errors']))

    def test_docx_same_quote_elsewhere_cannot_hide_wrong_verse(self):
        v2 = copy.deepcopy(self.study['verses'][0]); v2['key'] = '91:2'; v2['heading'] = 'الْآيَةُ الثَّانِيَةُ'
        v2['text'] = 'وَالْقَمَرِ إِذَا تَلَاهَا'
        for field in ('groups', 'themes', 'benefits'):
            for obj in v2[field]:
                obj['id'] += 'second'
        v2['groups'][0]['quotes'][0]['id'] = 'qsecond'
        v2['benefits'][0]['evidence_ids'] = ['qsecond']
        self.study['verses'].append(v2)
        file = self.make_docx()
        self.replace_docx_text(file, 'أَظْهَرَ ضَوْءَهَا.', '', count=1)
        self.assertIn('docx_wrong_verse_or_changed', self.codes(file))

    def test_quote_moved_under_wrong_axis_fails(self):
        from docx import Document
        file = self.make_docx(); doc = Document(file)
        quote = next(p for p in doc.paragraphs if p.style.name == 'Quotation')
        axis3 = next(p for p in doc.paragraphs if p.text.startswith('٣  '))
        quote._p.getparent().remove(quote._p); axis3._p.addnext(quote._p)
        doc.save(file)
        self.assertIn('docx_wrong_axis', self.codes(file))

    def test_header_vowels_are_checked(self):
        from docx import Document
        file = self.make_docx(); doc = Document(file)
        doc.sections[0].header.tables[0].cell(0,1).paragraphs[0].text = 'دراسة'
        doc.save(file)
        self.assertIn('docx_header_changed', self.codes(file))

    def test_multiple_surahs_keep_independent_seven_axes_in_one_word(self):
        from prepare_quran import prepare
        from build_docx import ASSETS
        self.study['surahs'] = [self.study.pop('surah'), {'number':92,'name':'اللَّيْلُ'}]
        second = copy.deepcopy(self.study['verses'][0])
        second['key'] = '92:1'; second['heading'] = 'سُورَةُ اللَّيْلِ — الْآيَةُ ١'
        q = prepare(ASSETS/'quran-simple-1.1.txt',json.loads((ASSETS/'quran-source.json').read_text(encoding='utf-8')))
        second['text'] = q['92:1']; self.quran['92:1'] = q['92:1']
        second['groups'][0]['id'] = 'gsecond'; second['groups'][0]['quotes'][0]['id'] = 'qsecond'
        second['themes'][0]['id'] = 'tsecond'; second['benefits'][0]['id'] = 'fsecond'
        second['benefits'][0]['evidence_ids'] = ['qsecond']; self.study['verses'].append(second)
        for source in self.sources['sources']:
            source['blocks'][0]['verse_keys'].append('92:1')
        self.coverage['items'].extend([{'id':'c3','source_id':'s01','block_id':'b001','verse_keys':['92:1'],'kind':'summary','status':'included','output_ids':['92:1']},{'id':'c4','source_id':'s02','block_id':'b001','verse_keys':['92:1'],'kind':'opinion','status':'included','output_ids':['qsecond']}])
        self.assertTrue(self.audit(self.make_docx())['passed'])

    def test_docx_term_definition_explanation_and_example_are_checked(self):
        file = self.make_docx()
        self.study["verses"][0]["terms"] = [{"id": "term1", "name": "مَجَازٌ", "field": "بَلَاغَةٌ",
                                            "definition": "تَعْرِيفُ الْمُصْطَلَحِ.", "explanation": "شَرْحُهُ هُنَا.",
                                            "example": "مِثَالٌ خَارِجِيٌّ.", "source_quote_ids": ["q1"]}]
        errors = self.audit(file)["errors"]
        missing = {e["path"] for e in errors if e["code"] == "docx_missing_text"}
        for field in ("definition", "explanation", "example"):
            self.assertIn(f"study.verses[0].terms[0].{field}", missing)

    def test_docx_preference_reason_cannot_disappear(self):
        file = self.make_docx()
        self.study["verses"][0]["groups"][0]["preferences"] = [{"source_id": "s02", "text": "هَذَا أَرْجَحُ.",
                                                               "reason": "لِدَلَالَةِ السِّيَاقِ.", "quote_ids": ["q1"]}]
        errors = self.audit(file)["errors"]
        self.assertTrue(any(e["code"] == "docx_missing_text" and e["path"].endswith(".reason") for e in errors))

    def test_cli_failure_is_nonzero_and_writes_report(self):
        del self.study["verses"][0]["themes"][0]["hadith"][0]["grade"]
        objects = {"study": self.study, "sources": self.sources,
                   "quran": self.quran, "coverage": self.coverage}
        args = [sys.executable, str(Path(__file__).with_name("audit_study.py"))]
        for name, value in objects.items():
            file = self.folder / f"{name}.json"
            file.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
            args += [f"--{name}", str(file)]
        out = self.folder / "report.json"
        result = subprocess.run(args + ["--out", str(out)], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(json.loads(out.read_text(encoding="utf-8"))["passed"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
