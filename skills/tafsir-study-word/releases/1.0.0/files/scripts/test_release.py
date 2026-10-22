"""Exercise release failures with a synthetic review record, never a real attestation."""
import copy
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx.oxml.ns import qn
from build_docx import Builder, ASSETS
from audit_layout import audit
from check_release import CHECKS, DOCUMENT_CHECKS, check, check_quote, digest
from test_tools import fixture


class ReleaseTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.study, self.sources, self.quran, self.coverage = fixture(self.root)
        self.study['verses'][0]['names']['kind'] = 'none'
        for item in self.coverage['items']:
            source = next(s for s in self.sources['sources'] if s['id'] == item['source_id'])
            block = source['blocks'][0]
            item.update(meaning='Synthetic test meaning', source_spans=[{'start':block['start'],'end':block['end']}])
        self.files = {}
        for name, obj in (('study', self.study), ('sources', self.sources), ('coverage', self.coverage), ('quran', self.quran)):
            path = self.root / (name + '.json'); path.write_text(json.dumps(obj), encoding='utf-8'); self.files[name] = path
        self.spec = json.loads((ASSETS / 'style-spec.json').read_text(encoding='utf-8'))
        self.files['docx'] = self.root / 'test.docx'; Builder(self.study, self.spec).content().save(self.files['docx'])
        from reportlab.pdfgen import canvas
        self.files['pdf'] = self.root / 'test.pdf'
        c = canvas.Canvas(str(self.files['pdf'])); c.drawString(50,750,'Synthetic release-gate fixture'); c.showPage(); c.save()
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(self.files['pdf'])); page = pdf[0]; bitmap = page.render(scale=1)
        self.png = self.root / 'page-1.png'; bitmap.to_pil().save(self.png); bitmap.close(); page.close(); pdf.close()
        self.review = {'schema_version':1,'reviewer':'Synthetic test only','reviewed_at':'2026-10-09','input_sha256':{k:digest(v) for k,v in self.files.items()},'open_issues':[], 'source_reading':[], 'verses':[{'key':'91:1','checks':{k:{'status':'pass','note':'Synthetic record used solely to test validation.'} for k in CHECKS}}], 'verification':[], 'visual':{'engine':'Synthetic test','fields_updated':True,'pages':[{'number':1,'path':str(self.png),'sha256':digest(self.png),'status':'pass','note':'Synthetic fixture; not an assertion of real visual review.'}]}}
        for s in self.sources['sources']:
            for b in s['blocks']:
                self.review['source_reading'].append({'source_id':s['id'],'block_id':b['id'],'start':b['start'],'end':b['end'],'status':'read_and_inventoried','note':'Synthetic fixture.'})
        self.review['document_checks'] = {k:{'status':'pass','note':'Synthetic test only.'} for k in DOCUMENT_CHECKS}
        for genre in ('hadith','poetry'):
            self.review['verification'].append({'path':f'study.verses[0].themes[0].{genre}[0]','status':'verified','url':'https://example.com','locator':'fixture','checked_at':'2026-10-09','finding':'Synthetic fixture.'})

    def tearDown(self):
        self.tmp.cleanup()

    def result(self):
        return check(self.study, self.sources, self.coverage, self.review, self.files)

    def codes(self):
        return {e['code'] for e in self.result()['errors']}

    def test_complete_synthetic_record(self):
        self.assertTrue(self.result()['passed'])

    def test_changed_docx_invalidates_review(self):
        self.files['docx'].write_bytes(self.files['docx'].read_bytes() + b'changed')
        self.assertIn('stale_review', self.codes())

    def test_missing_linguistic_review_blocks_release(self):
        del self.review['verses'][0]['checks']['vocalization']
        self.assertIn('scientific_review_missing', self.codes())

    def test_front_matter_review_is_required(self):
        del self.review['document_checks']['front_matter_vocalization']
        self.assertIn('document_review_missing', self.codes())

    def test_unread_source_block_blocks_release(self):
        self.review['source_reading'].pop()
        self.assertIn('unread_source_block', self.codes())

    def test_partial_source_reading_cannot_pass(self):
        self.review['source_reading'][0]['end'] -= 1
        self.assertIn('reading_extent', self.codes())

    def test_unverified_hadith_blocks_release(self):
        self.review['verification'].pop(0)
        self.assertIn('unverified_evidence', self.codes())

    def test_name_without_alnahj_reference_blocks_release(self):
        self.study['verses'][0]['names']['kind'] = 'name'
        self.assertIn('alnahj_missing', self.codes())

    def test_coarse_coverage_without_spans_blocks_release(self):
        self.coverage['items'][0].pop('source_spans')
        self.assertIn('coverage_not_inventory', self.codes())

    def test_unviewed_page_blocks_release(self):
        self.review['visual']['pages'] = []
        self.assertIn('unreviewed_pages', self.codes())

    def test_changed_page_image_invalidates_review(self):
        self.png.write_bytes(self.png.read_bytes()+b'changed')
        self.assertIn('stale_page_image', self.codes())

    def test_deletion_of_meaning_is_not_an_approved_omission(self):
        errors = []
        q={'source_spans':[{'start':0,'end':5},{'start':8,'end':11}],'omissions':[{'start':5,'end':8,'kind':'shortening','reason':'shorter'}]}
        check_quote(q,'quote',lambda *e:errors.append(e))
        self.assertEqual(errors[0][0], 'unapproved_omission')

    def test_summary_selection_is_limited_to_extracted_summary(self):
        errors=[]
        q={'mode':'extracted','source_spans':[{'start':0,'end':5},{'start':8,'end':11}],'omissions':[{'start':5,'end':8,'kind':'summary_selection','reason':'Selection for axis-one summary.'}]}
        check_quote(q,'summary',lambda *e:errors.append(e),summary=True)
        self.assertEqual(errors,[])
        check_quote(q,'opinion',lambda *e:errors.append(e))
        self.assertEqual(errors[0][0],'unapproved_omission')

    def test_real_builder_layout_passes(self):
        self.assertTrue(audit(self.files['docx'], self.spec)['passed'])

    def test_wrong_margins_fail_layout(self):
        doc = Builder(self.study,self.spec).content(); doc.sections[0].left_margin = 100
        doc.save(self.files['docx'])
        self.assertIn('margins', audit(self.files['docx'],self.spec)['issues'])

    def test_missing_corner_fails_layout(self):
        doc = Builder(self.study,self.spec).content()
        run = doc.sections[0].header.tables[0].cell(0,0).paragraphs[0].runs[0]
        run.clear(); doc.save(self.files['docx'])
        self.assertIn('four_corner_ornaments', audit(self.files['docx'],self.spec)['issues'])

    def test_reference_asset_is_never_modified(self):
        before = digest(ASSETS/'reference.docx')
        Builder(self.study,self.spec).content().save(self.files['docx'])
        self.assertEqual(before,digest(ASSETS/'reference.docx'))


if __name__ == '__main__':
    unittest.main(verbosity=2)
