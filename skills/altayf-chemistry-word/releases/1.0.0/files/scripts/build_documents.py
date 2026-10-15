"""Build the five editable documents from a reviewed lesson-plan.json. Never invents lesson text."""
import argparse, re, shutil
from copy import deepcopy
from pathlib import Path
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from common import ROOT, NS, read_json, write_json, digest, val, font, paragraph_settings, text_runs, canonical_properties
from header_badge import set_title_and_unit

COLORS={'definition':'94CB7A','law':'B59AC9','example':'F6A34C','exercises':'69C8C9',
        'practice':'69C8C9','note':'FFC450','prestart':'A6CDD4','sources':'C3D5E6','question':'69C8C9','rule':'B59AC9'}
NAMES=['01-الدرس-المنسق.docx','02-حلول-الدرس-للمعلم.docx','03-المكثف-وبنك-الأسئلة.docx',
       '04-حلول-المكثف-والبنك.docx','05-الشرح-المعاد-صياغته-والمنسق.docx']
class Builder:
    def __init__(self,plan,plan_dir,out,inventory):
        self.plan=plan;self.base=plan_dir;self.out=out;self.inventory=inventory
        self.questions={q['id']:q for q in plan['questions']};self.math={}
        invdir=Path(plan['inventory']).resolve() if Path(plan['inventory']).is_absolute() else (plan_dir/plan['inventory']).resolve()
        for m in inventory['equations']:self.math[m['id']]=E.fromstring((invdir.parent/m['file']).read_bytes())
        for mid,xml in plan.get('new_equations',{}).items():
            n=E.fromstring(xml.encode('utf-8'),E.XMLParser(resolve_entities=False,no_network=True))
            if n.tag not in [val for val in ['{'+NS['m']+'}oMath','{'+NS['m']+'}oMathPara']]:raise ValueError('Equation must be native OMML')
            self.math[mid]=n
    def path(self,p):return Path(p).resolve() if Path(p).is_absolute() else (self.base/p).resolve()
    def mathpara(self,container,mid,size=12.5,color='182742',bold=False):
        if mid not in self.math:raise ValueError('Unknown math '+mid)
        p=self.para(container,'',center=True,size=size)
        node=deepcopy(self.math[mid])
        for mr in node.findall('.//m:r',NS):
            pr=mr.find('w:rPr',NS)
            if pr is None:
                pr=E.Element('{'+NS['w']+'}rPr')
                mr.insert(1 if mr.find('m:rPr',NS) is not None else 0,pr)
            val(pr,'sz',round(size*2));val(pr,'szCs',round(size*2))
            val(pr,'b',int(bold));val(pr,'bCs',int(bold))
            val(pr,'color',color)
        p._p.append(node)
        return p
    def para(self,container,value='',size=None,bold=False,center=False,style=None,color=None):
        p=container.add_paragraph(style=style);paragraph_settings(p,center=center,compact=self.compact)
        kind=getattr(self,'current_kind',None)
        default=(11 if kind=='practice' else 10.5 if kind=='exercises' else 12) if self.compact else (11 if kind=='practice' else 10.5 if kind=='exercises' else 13 if kind in ['law','definition','note'] else 14)
        text_runs(p,value,size or default,bold,color or ('082C5A' if bold else '182742'),self.math)
        return p
    def numbering(self,ordered):
        root=self.doc.part.numbering_part.element
        aid=max([int(n.get('{'+NS['w']+'}abstractNumId')) for n in root.findall('w:abstractNum',NS)]+[0])+1
        nid=max([int(n.get('{'+NS['w']+'}numId')) for n in root.findall('w:num',NS)]+[0])+1
        an=E.Element('{'+NS['w']+'}abstractNum',{'{'+NS['w']+'}abstractNumId':str(aid)})
        first_num=root.find('w:num',NS)
        root.insert(root.index(first_num) if first_num is not None else len(root),an)
        lv=E.SubElement(an,'{'+NS['w']+'}lvl',{'{'+NS['w']+'}ilvl':'0'})
        for name,value in [('start','1'),('numFmt','decimal' if ordered else 'bullet'),('lvlText','%1.' if ordered else '•'),('lvlJc','left')]:val(lv,name,value)
        pr=E.SubElement(lv,'{'+NS['w']+'}pPr');ind=E.SubElement(pr,'{'+NS['w']+'}ind');ind.set('{'+NS['w']+'}left','480');ind.set('{'+NS['w']+'}hanging','260')
        rp=E.SubElement(lv,'{'+NS['w']+'}rPr');f=E.SubElement(rp,'{'+NS['w']+'}rFonts');f.set('{'+NS['w']+'}ascii','Dubai');f.set('{'+NS['w']+'}hAnsi','Dubai');f.set('{'+NS['w']+'}cs','Dubai')
        num=E.SubElement(root,'{'+NS['w']+'}num',{'{'+NS['w']+'}numId':str(nid)});val(num,'abstractNumId',aid)
        ov=E.SubElement(num,'{'+NS['w']+'}lvlOverride',{'{'+NS['w']+'}ilvl':'0'});val(ov,'startOverride','1')
        return nid
    def table(self,container,headers,rows,widths=None,size=10.5,color='BFD1E2'):
        cols=len(headers);widths=widths or [16.1/cols]*cols
        if len(widths)!=cols or sum(widths)>17.6:raise ValueError('Invalid table width')
        t=container.add_table(rows=1,cols=cols);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        val(t._tbl.tblPr,'bidiVisual',1)
        val(t._tbl.tblPr,'tblCaption','AltayfDataTable')
        for n,w in zip(t._tbl.tblGrid,widths):n.set('{'+NS['w']+'}w',str(Cm(w).twips))
        tw=t._tbl.tblPr.find('w:tblW',NS);tw.set('{'+NS['w']+'}type','dxa');tw.set('{'+NS['w']+'}w',str(Cm(sum(widths)).twips))
        for i,values in enumerate([headers]+rows):
            if len(values)!=cols:raise ValueError('Wrong number of cells')
            row=t.rows[0] if i==0 else t.add_row();val(row._tr.get_or_add_trPr(),'cantSplit',1)
            if i==0:val(row._tr.get_or_add_trPr(),'tblHeader',1)
            for j,value in enumerate(values):
                c=row.cells[j];c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=c._tc.get_or_add_tcPr();mar=E.SubElement(pr,'{'+NS['w']+'}tcMar')
                for side,v in [('top',100),('bottom',100),('left',120),('right',120)]:
                    n=E.SubElement(mar,'{'+NS['w']+'}'+side);n.set('{'+NS['w']+'}w',str(v));n.set('{'+NS['w']+'}type','dxa')
                br=E.SubElement(pr,'{'+NS['w']+'}tcBorders')
                for side in ['top','bottom','left','right']:
                    n=E.SubElement(br,'{'+NS['w']+'}'+side);n.set('{'+NS['w']+'}val','single');n.set('{'+NS['w']+'}sz','5');n.set('{'+NS['w']+'}color',color)
                p=c.paragraphs[0];paragraph_settings(p,center=i==0 or j in [0,cols-1],compact=True)
                if i==0:p.paragraph_format.keep_with_next=True
                if i>0 and headers[0]=='الرمز' and j in [0,cols-1] and not isinstance(value,list):value=[{'text':str(value),'ltr':True}]
                text_runs(p,value,size,i==0 or j==0,'082C5A' if i==0 else '182742',self.math)
        return t
    def ribbon(self,c,kind,number=None,paragraph=None):
        if number is None:
            self.count[kind]=self.count.get(kind,0)+1;number=self.count[kind]
        filename=f'label-{kind}-{number:02d}.png'
        p=self.out/'images'/'ribbons'/filename
        if not p.exists():
            asset=ROOT/'assets'/'ribbons'/filename
            if not asset.exists() and kind=='rule':asset=ROOT/'assets'/'rule-references'/f'rule-{number}.png'
            if asset.exists():p.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(asset,p)
        if not p.exists():raise ValueError(f'Missing approved/generated numbered ribbon {filename}; generate before build.')
        target=paragraph if paragraph is not None else self.para(c,'',size=1)
        target.add_run().add_picture(str(p),width=Cm(4.0))
        return number
    def card(self,container,kind,blocks,qid=None,show_ribbon=True,number=None):
        if kind not in COLORS:raise ValueError('Unknown card kind '+kind)
        t=container.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Cm(17.6)
        t.alignment=WD_TABLE_ALIGNMENT.CENTER;c=t.cell(0,0);c.width=Cm(17.6)
        # Deliberately unique marker color, replaced after PDF geometry measurement.
        marker=f'{0xE10000+len(self.cards)+1:06X}';name=f'AltayfCard{len(self.cards)+1:05d}'
        pr=c._tc.get_or_add_tcPr();br=E.SubElement(pr,'{'+NS['w']+'}tcBorders')
        for side in ['top','bottom','left','right']:
            n=E.SubElement(br,'{'+NS['w']+'}'+side);n.set('{'+NS['w']+'}val','single');n.set('{'+NS['w']+'}sz','5');n.set('{'+NS['w']+'}color',marker)
        mar=E.SubElement(pr,'{'+NS['w']+'}tcMar')
        for side in ['top','bottom','left','right']:
            n=E.SubElement(mar,'{'+NS['w']+'}'+side);n.set('{'+NS['w']+'}w','180');n.set('{'+NS['w']+'}type','dxa')
        p=c.paragraphs[0];paragraph_settings(p,compact=self.compact)
        p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
        bid=str(10000+len(self.cards));bs=E.SubElement(p._p,'{'+NS['w']+'}bookmarkStart',{'{'+NS['w']+'}id':bid,'{'+NS['w']+'}name':name});E.SubElement(p._p,'{'+NS['w']+'}bookmarkEnd',{'{'+NS['w']+'}id':bid})
        reuse=self.reference_ribbons.get(qid) if isinstance(qid,str) else None
        key=(kind,reuse['number']) if reuse else None
        if not show_ribbon or (self.teacher and key is not None and key==self.last_teacher_ribbon):
            ribbon_added=False
            number=number if number is not None else (reuse['number'] if reuse else self.count.get(kind,1))
        else:
            ribbon_added=True
            number=self.ribbon(c,kind,reuse['number'] if reuse else number,paragraph=p)
            # A fixed one-point line clips the inline PNG in Word even though
            # the image remains present in OOXML. Let this line grow to image
            # height; ribbonless cards keep the tiny bookmark line.
            p.paragraph_format.line_spacing=1.07
            p.paragraph_format.space_after=Pt(4)
        if self.teacher:self.last_teacher_ribbon=key
        for key in (qid if isinstance(qid,list) else [qid]):
            if key:self.ribbon_map[key]={'kind':kind,'number':number}
        previous=getattr(self,'current_kind',None);self.current_kind=kind
        self.blocks(c,blocks);self.current_kind=previous
        # Small cards stay whole; long content must be split intentionally and reviewed.
        ps=c._tc.findall('./w:p',NS)
        for pnode in ps[:-1]:val(pnode.get_or_add_pPr(),'keepNext',1)
        self.cards.append({'name':name,'marker':marker,'border':COLORS[kind],'tint':COLORS[kind],'kind':kind,'question_id':qid,'ribbon_number':number,'visible_ribbon_expected':ribbon_added})
        self.para(container,'',size=1).paragraph_format.space_after=Pt(4)
        return number
    def answer_area(self,container,lines=1):
        """Separate, transparent writing space below an unsolved training/exercise card."""
        self.answer_areas+=1
        t=container.add_table(rows=1,cols=1);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t.columns[0].width=Cm(17.0);c=t.cell(0,0);c.width=Cm(17.0)
        val(t._tbl.tblPr,'tblCaption','AltayfAnswerArea')
        pr=c._tc.get_or_add_tcPr();br=E.SubElement(pr,'{'+NS['w']+'}tcBorders')
        for side in ['top','bottom','left','right']:
            n=E.SubElement(br,'{'+NS['w']+'}'+side);n.set('{'+NS['w']+'}val','single');n.set('{'+NS['w']+'}sz','4');n.set('{'+NS['w']+'}color','A7C7D8')
        for i in range(max(1,min(int(lines),5))):
            p=c.paragraphs[0] if i==0 else c.add_paragraph()
            paragraph_settings(p,compact=True);p.paragraph_format.space_after=Pt(1)
            text_runs(p,'....................................................................................................................................',8,False,'9FB8C9')
        self.para(container,'',size=1).paragraph_format.space_after=Pt(3)
    def worked_solution(self,container,q):
        w=q['worked_solution']
        self.para(container,[{'text':'المسألة:','bold':True,'color':'B8621A'}],size=12.5)
        self.para(container,q.get('prompt_runs',q['prompt']),size=12.5)
        label=self.para(container,[{'text':'1. المعطيات:','bold':True,'color':'0B3B67'}],size=12.5)
        label.paragraph_format.keep_with_next=True
        if w.get('givens_math_id'):self.mathpara(container,w['givens_math_id'],size=12.5,color='0B3B67')
        else:self.para(container,w['givens'],size=12.5)
        self.para(container,[{'text':'2. المطلوب:','bold':True,'color':'157A70'}, {'text':' '+w['target']}],size=12.5)
        for i,step in enumerate(w['steps'],3):
            color=step.get('color', '6B4E93' if 'قانون' in step['label'] else 'B8621A')
            p=self.para(container,[{'text':str(i)+'. '+step['label'],'bold':True,'color':color}],size=12.5)
            p.paragraph_format.keep_with_next=True
            self.mathpara(container,step['math_id'],size=12.5,color=color,bold=False)
            if step.get('note_runs') or step.get('note'):
                self.para(container,step.get('note_runs',step.get('note')),size=12.5)
        self.para(container,[{'text':'الإجابة','bold':True,'color':'A8500B'}],size=10,center=False)
        self.mathpara(container,w['final_math_id'],size=13,color='B8621A',bold=True)
        if w.get('conclusion_runs') or w.get('conclusion'):
            self.para(container,w.get('conclusion_runs',w.get('conclusion')),size=12.5)
    def glossary(self,container,entries):
        for entry in entries:
            t=container.add_table(rows=1,cols=1);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
            t.columns[0].width=Cm(17.0);c=t.cell(0,0);c.width=Cm(17.0)
            val(t._tbl.tblPr,'tblCaption','AltayfGlossary')
            pr=c._tc.get_or_add_tcPr();br=E.SubElement(pr,'{'+NS['w']+'}tcBorders')
            for side in ['top','bottom','left','right']:
                n=E.SubElement(br,'{'+NS['w']+'}'+side);n.set('{'+NS['w']+'}val','single');n.set('{'+NS['w']+'}sz','4');n.set('{'+NS['w']+'}color','BFD1E2')
            p=c.paragraphs[0];paragraph_settings(p,compact=True);text_runs(p,entry['term'],11,True,'082C5A',self.math)
            if entry.get('definition') or entry.get('definition_runs'):
                p=c.add_paragraph();paragraph_settings(p,compact=True)
                text_runs(p,entry.get('definition_runs',entry.get('definition')),10.5,False,'182742',self.math)
            self.para(container,'',size=1).paragraph_format.space_after=Pt(4)
    def remember(self,container,items):
        self.para(container,'تذكّر عند الحل',size=15,bold=True).paragraph_format.keep_with_next=True
        nid=self.numbering(False)
        for item in items:
            p=self.para(container,item.get('runs',item.get('text')) if isinstance(item,dict) else item,size=14)
            pr=p._p.get_or_add_pPr();np=E.SubElement(pr,'{'+NS['w']+'}numPr');val(np,'ilvl',0);val(np,'numId',nid)
    def question(self,container,qid,worked=False):
        q=self.questions[qid];self.qids.append(qid)
        if self.teacher:
            role='compact' if self.teacher_compact else 'student'
            locator=self.para(container,f"{q.get('topic','')} | {q['id']} | موضعه في ملف الطالب: {q.get('locators',{}).get(role,'غير مسجل')}",size=11,bold=True)
            locator.paragraph_format.keep_with_next=True
        solved=worked or self.teacher
        blocks=[]
        if solved and q.get('worked_solution'):
            blocks=[{'type':'worked_solution','question_id':qid}]
        else:
            blocks=[{'type':'paragraph','runs':q.get('prompt_runs',q['prompt']),'size':10.5 if q.get('ribbon')=='exercises' else 11 if q.get('ribbon')=='practice' else 11.5 if self.compact else 14}]
        blocks.extend(q.get('question_blocks',[]))
        for o in q.get('options',[]):blocks.append({'type':'paragraph','text':f"{o['id']}) {o['text']}"})
        if solved and not q.get('worked_solution'):blocks.extend(q['solution'])
        if self.teacher and qid in self.reference_ribbons:
            kind=self.reference_ribbons[qid]['kind']
        elif self.compact and q['kind']=='example' and not worked:
            # A compact question bank carries no worked examples. Reuse the
            # source example as an unsolved practice with a writing area.
            kind='practice'
        else:
            kind='example' if q['kind']=='example' else q.get('ribbon','practice')
        self.card(container,kind,blocks,qid)
        if not solved and kind in ('practice','exercises','question'):
            self.answer_area(container,q.get('answer_lines',1))
    def blocks(self,container,blocks):
        for b in blocks:
            typ=b['type']
            if typ=='heading':
                p=self.para(container,b['text'],size=17 if b.get('level',2)==1 else 15,bold=True,style='Heading '+str(b.get('level',2)))
                p.paragraph_format.keep_with_next=True
            elif typ=='paragraph':self.para(container,b.get('runs',b.get('text','')),size=b.get('size'),bold=b.get('bold',False))
            elif typ=='list':
                nid=self.numbering(b.get('ordered',False))
                for item in b['items']:
                    value=item.get('runs',item.get('text','')) if isinstance(item,dict) else item
                    p=self.para(container,value);pr=p._p.get_or_add_pPr();np=E.SubElement(pr,'{'+NS['w']+'}numPr');val(np,'ilvl',0);val(np,'numId',nid)
            elif typ=='math':
                self.mathpara(container,b['math_id'],size=b.get('size',13),color=b.get('color','182742'),bold=b.get('bold',False))
            elif typ=='worked_solution':self.worked_solution(container,self.questions[b['question_id']])
            elif typ=='definition':
                inner=[{'type':'paragraph','runs':[{'text':b['term'],'bold':True,'color':'082C5A'}],'size':13},
                       {'type':'paragraph','runs':b.get('definition_runs',b.get('definition','')),'size':13}]
                self.card(container,'definition',inner)
            elif typ=='law':
                inner=[]
                if b.get('name'):inner.append({'type':'paragraph','text':b['name'],'bold':True,'size':13})
                inner.extend(b.get('blocks',[]));inner.extend({'type':'math','math_id':x,'color':'6B4E93','size':13,'bold':True} for x in b['math_ids'])
                inner.append({'type':'paragraph','text':'الرموز ومعانيها ووحداتها','size':10.5,'bold':True})
                inner.append({'type':'table','headers':['الرمز','المعنى','الوحدة'],'rows':b['symbols'],'widths':[2.3,9.5,4.3],'size':10.5,'color':'D4C2E2'})
                self.card(container,'law',inner)
            elif typ in ['example','question']:self.question(container,b['question_id'],worked=typ=='example' and not self.compact)
            elif typ=='question_group':
                # One numbered banner at the start; separate prompt and writing
                # area for each item, matching the approved lesson edition.
                kind=b.get('ribbon','exercises');number=None
                for i,qid in enumerate(b['question_ids'],b.get('start',1)):
                    q=self.questions[qid];self.qids.append(qid)
                    number_label=str(i)+(f' | {q["difficulty_label"]}' if q.get('difficulty_label') else '')
                    inner=[{'type':'paragraph','runs':[{'text':number_label,'bold':True,'color':'082C5A'}], 'size':10.5},
                           {'type':'paragraph','runs':q.get('prompt_runs',q['prompt']),'size':10.5}]
                    inner.extend(q.get('question_blocks',[]))
                    for o in q.get('options',[]):inner.append({'type':'paragraph','text':f"{o['id']}) {o['text']}"})
                    if self.teacher:
                        inner.extend([{'type':'worked_solution','question_id':qid}] if q.get('worked_solution') else q['solution'])
                    number=self.card(container,kind,inner,qid,show_ribbon=(i==b.get('start',1)),number=number)
                    if not self.teacher:self.answer_area(container,q.get('answer_lines',1))
            elif typ=='card':self.card(container,b['kind'],b['blocks'])
            elif typ=='glossary':self.glossary(container,b['entries'])
            elif typ=='final_summary':
                self.para(container,b.get('title','ملخص القوانين | مراجعة نهائية'),size=15,bold=True).paragraph_format.keep_with_next=True
                self.blocks(container,b['blocks'])
            elif typ=='remember_at_solution':self.remember(container,b['items'])
            elif typ=='table' and b.get('section')=='glossary':
                self.glossary(container,[{'term':r[0],'definition':r[1] if len(r)>1 else ''} for r in b['rows']])
            elif typ=='table':self.table(container,b['headers'],b['rows'],b.get('widths'),b.get('size',10.5),b.get('color','BFD1E2'))
            elif typ=='figure':
                image=self.path(b['path'])
                copied=self.out/'images'/'science'/(b.get('image_id',image.stem)+image.suffix)
                copied.parent.mkdir(parents=True,exist_ok=True)
                if copied.resolve()!=image.resolve():shutil.copy2(image,copied)
                image=copied;p=self.para(container,'',center=True)
                p.add_run().add_picture(str(image),width=Cm(b.get('width_cm',12.8)))
                if b.get('caption'):self.para(container,b['caption'],size=12,center=True)
            else:raise ValueError('Unsupported block '+typ+'; extend deliberately, never silently drop it.')
    def build(self,role,blocks,filename,teacher=False,compact=False,reference_ribbons=None):
        self.doc=Document(ROOT/'assets'/'templates'/'lesson-frame.docx')
        set_title_and_unit(self.doc,self.plan['lesson_title'],self.plan['unit_number'])
        self.teacher=teacher;self.teacher_compact=role=='compact_teacher';self.compact=compact;self.count={};self.cards=[];self.qids=[];self.answer_areas=0;self.current_kind=None;self.ribbon_map={};self.reference_ribbons=reference_ribbons or {};self.last_teacher_ribbon=None
        if teacher:self.para(self.doc,self.plan['lesson_title']+' '+{'teacher':'حلول الدرس للمعلم','compact_teacher':'حلول المكثف والبنك'}[role],size=18,bold=True,style='Title')
        self.blocks(self.doc,blocks);path=self.out/filename
        canonical_properties(self.doc.element);canonical_properties(self.doc.part.numbering_part.element)
        self.doc.save(path)
        write_json(self.out/'qa'/f'{role}-cards.json',self.cards)
        write_json(self.out/'qa'/f'{role}-questions.json',self.qids)
        return {'role':role,'path':str(path),'cards':str(self.out/'qa'/f'{role}-cards.json'),'question_ids':self.qids,'answer_areas':self.answer_areas,'ribbon_map':self.ribbon_map,'sha256_before_render':digest(path)}
def main():
    ap=argparse.ArgumentParser();ap.add_argument('plan');ap.add_argument('outdir');a=ap.parse_args();pp=Path(a.plan).resolve();p=read_json(pp)
    out=Path(a.outdir).resolve();out.mkdir(parents=True,exist_ok=True)
    from validate_plan import validate
    validation=validate(p,pp.parent)
    write_json(out/'qa'/'plan-build-gate.json',validation)
    if validation['errors']:raise ValueError('Content plan failed validation: '+'; '.join(validation['errors']))
    background=out/'images'/'background'/'background.png';background.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(ROOT/'assets'/'page'/'background.png',background)
    inv=Path(p['inventory']);inv=inv if inv.is_absolute() else pp.parent/inv
    inventory=read_json(inv)
    if p['source_sha256']!=inventory['sha256'] or digest(inventory['source'])!=inventory['sha256']:raise ValueError('Source changed. Re-inventory and review.')
    b=Builder(p,pp.parent,out,inventory);manifest=[]
    for role,key,name in [('student','student',NAMES[0]),('teacher','student',NAMES[1]),('compact','compact',NAMES[2]),('compact_teacher','compact',NAMES[3]),('rewritten','rewritten',NAMES[4])]:
        if 'teacher' in role:
            # Same order and exactly the same question IDs as the corresponding student file.
            original=manifest[0 if role=='teacher' else 2]['question_ids']
            blocks=[{'type':'question','question_id':q} for q in original]
            source_blocks=p['documents']['student' if role=='teacher' else 'compact']
            blocks.extend(b for b in source_blocks if b['type']=='remember_at_solution')
        else:blocks=p['documents'][key]
        reference_ribbons=manifest[0 if role=='teacher' else 2]['ribbon_map'] if 'teacher' in role else None
        manifest.append(b.build(role,blocks,name,teacher='teacher' in role,compact=role=='compact',reference_ribbons=reference_ribbons))
    write_json(out/'output-manifest.json',manifest);print('Built five editable DOCX files. Rendering, geometry/tint pass and visual inspection are still required.')
if __name__=='__main__':main()
