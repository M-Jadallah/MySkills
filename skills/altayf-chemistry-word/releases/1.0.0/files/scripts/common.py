"""Shared, local-only JSON and Word primitives. Requires python-docx, lxml, Pillow."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import json, re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'mc':'http://schemas.openxmlformats.org/markup-compatibility/2006',
      'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
def read_json(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def write_json(p, value):
    p=Path(p);p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
def digest(p): return sha256(Path(p).read_bytes()).hexdigest()
def norm(s): return re.sub(r'[\s\u200e\u200f\u202a-\u202e\u2066-\u2069]+',' ',s).strip()
def val(parent, name, value):
    n=parent.find(qn('w:'+name))
    if n is None:n=OxmlElement('w:'+name);parent.append(n)
    n.set(qn('w:val'),str(value));return n
def font(r,size=14,bold=False,color='182742',ltr=False):
    r.font.name='Dubai';r.font.size=Pt(size);r.bold=bold
    r.font.color.rgb=RGBColor.from_string(color)
    pr=r._r.get_or_add_rPr();val(pr,'szCs',round(size*2));val(pr,'bCs',int(bold))
    rf=pr.find(qn('w:rFonts'))
    for k in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+k),'Dubai')
    val(pr,'rtl',0 if ltr else 1)
def paragraph_settings(p,size=14,center=False,compact=False):
    pr=p._p.get_or_add_pPr();val(pr,'bidi',1)
    # In this retained bidi template default alignment is right. w:jc right mirrors it.
    old=pr.find(qn('w:jc'))
    if old is not None:pr.remove(old)
    if center:val(pr,'jc','center')
    p.paragraph_format.space_after=Pt(4 if compact else 7)
    p.paragraph_format.line_spacing=1.07 if compact else 1.16
    p.paragraph_format.widow_control=True
def text_runs(p, value, size=14, bold=False, color='182742', math=None):
    items=value if isinstance(value,list) else [value] if isinstance(value,dict) else [{'text':str(value)}]
    for it in items:
        if 'math_id' in it:
            if math is None or it['math_id'] not in math:raise ValueError('Unknown math_id '+it['math_id'])
            p._p.append(deepcopy(math[it['math_id']]));continue
        s=it.get('text','');strong=it.get('bold',bold);ink=it.get('color',color)
        # A source may glue Arabic prose to a quantity/unit. Add presentation
        # spacing without altering chemical formulas or Latin variable tokens.
        s=re.sub(r'([\u0621-\u065f])(?=[A-Za-z0-9α-ωΑ-Ω])',lambda m:m.group(1)+'\u00a0',s)
        s=re.sub(r'([A-Za-z0-9α-ωΑ-Ω₀-₉⁰-⁹])(?=[\u0621-\u065f])',lambda m:m.group(1)+'\u00a0',s)
        if 'ltr' in it:
            font(p.add_run(s),size,strong,ink,it['ltr']);continue
        # Keep paired parentheses, Latin, decimal numbers and units in coherent LTR runs.
        # A complete calculation must share one LTR run; splitting at '=' or
        # spaced operators can visually reverse a fraction in an RTL paragraph.
        chars=r'A-Za-z0-9α-ωΑ-Ω₀-₉⁰-⁹.,،؛/%°·×=+−–\-:()\[\] \t'
        pat=rf'(?=[{chars}]*[A-Za-z0-9α-ωΑ-Ω])[A-Za-z0-9α-ωΑ-Ω(\[][{chars}]*'
        last=0
        mixed=bool(re.search(r'[\u0621-\u065f]',s))
        for m in re.finditer(pat,s):
            if m.start()>last:
                segment=s[last:m.start()]
                if mixed:segment=segment.rstrip(' \t\u00a0').lstrip(' \t\u00a0') if last else segment.rstrip(' \t\u00a0')
                if segment:font(p.add_run(segment),size,strong,ink)
            # Word moves a boundary space with its run direction. Padding both
            # edges INSIDE a coherent LTR island keeps Arabic words separate
            # from quantities visually, including source text without spaces.
            token=m.group()
            if mixed:
                # Keep an entire parenthesized scientific tuple together. An
                # Arabic comma OUTSIDE brackets separates independent LTR
                # islands, while a comma inside a tuple stays with its values.
                parts=[];begin=0;depth=0
                for pos,ch in enumerate(token):
                    if ch in '([':depth+=1
                    elif ch in ')]':depth=max(0,depth-1)
                    elif ch in '،؛' and depth==0:
                        parts.append((token[begin:pos],True));parts.append((ch,False));begin=pos+1
                parts.append((token[begin:],True))
                for part,is_ltr in parts:
                    if not is_ltr:font(p.add_run(part),size,strong,ink);continue
                    part=part.strip(' \t\u00a0')
                    if not part:continue
                    punct=''
                    end=re.search(r'[.,:]+$',part)
                    if end:punct=end.group();part=part[:end.start()]
                    part=re.sub(r'\([^()]*\)',lambda g:g.group().replace(' ','\u00a0'),part)
                    part=re.sub(r'(?<=[0-9₀-₉]) +(?=[A-Za-zα-ωΑ-Ω°])','\u00a0',part)
                    if part:font(p.add_run('\u00a0'+part+'\u00a0'),size,strong,ink,True)
                    if punct:font(p.add_run(punct),size,strong,ink)
            else:font(p.add_run(token),size,strong,ink,True)
            last=m.end()
        if last<len(s):
            tail=s[last:].lstrip(' \t\u00a0') if mixed and last else s[last:]
            if tail:font(p.add_run(tail),size,strong,ink)
def iter_blocks(blocks):
    for b in blocks:
        yield b
        for key in ['blocks','items']:
            children=b.get(key,[])
            if children and isinstance(children[0],dict) and 'type' in children[0]:yield from iter_blocks(children)

def canonical_properties(root):
    """Use WML schema order; a tolerant Word open is not proof tagged PDF can consume it."""
    orders={
      'rPr':'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange',
      'pPr':'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
      'tblPr':'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription tblPrChange',
      'tcPr':'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark headers tcPrChange',
      'trPr':'cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden ins del trPrChange',
      'numbering':'numPicBullet abstractNum num numIdMacAtCleanup',
      'tcBorders':'top left start bottom right end insideH insideV tl2br tr2bl',
      'tblBorders':'top left start bottom right end insideH insideV',
      'tcMar':'top left start bottom right end',
      'tblCellMar':'top left start bottom right end'
    }
    for parent in root.iter():
        if not isinstance(parent.tag,str):continue
        local=parent.tag.split('}')[-1]
        if not parent.tag.startswith('{'+NS['w']+'}') or local not in orders:continue
        rank={name:i for i,name in enumerate(orders[local].split())}
        children=list(parent)
        sorted_nodes=sorted(children,key=lambda n:rank.get(n.tag.split('}')[-1],999))
        if children!=sorted_nodes:
            for n in children:parent.remove(n)
            for n in sorted_nodes:parent.append(n)

