from __future__ import annotations

from docx import Document
from docx.shared import Inches

from lesson_docx.builder import build_lesson_docx
from lesson_docx.tables import USABLE_WIDTH_INCHES
from lesson_docx.validate import validate_docx


def _sample_with_table(headers: list[str], rows: list[list[str]]) -> dict:
    return {
        "title": "عنوان",
        "subtitle": "",
        "sections": [
            {
                "heading": "محور",
                "blocks": [{"type": "paragraph", "text": "نص."}],
                "subsections": [],
                "section_summary": "",
                "table": {"title": "جدول", "headers": headers, "rows": rows},
                "concept_map": None,
            }
        ],
    }


def test_two_column_table_renders_correctly(tmp_path):
    out = tmp_path / "out.docx"
    data = _sample_with_table(["العمود الأول", "العمود الثاني"], [["أ", "ب"]])
    build_lesson_docx(data, out)

    doc = Document(str(out))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.columns) == 2
    assert len(table.rows) == 2  # header + 1 data row

    errors, _warnings = validate_docx(out)
    assert errors == []


def test_six_column_five_row_table_renders_without_truncation(tmp_path):
    out = tmp_path / "out.docx"
    letters = ["أ", "ب", "ج", "د", "هـ", "و"]
    headers = [f"عمود {letters[i]}" for i in range(6)]
    rows = [[f"صف{letters[r]}عمود{letters[c]}" for c in range(6)] for r in range(5)]
    data = _sample_with_table(headers, rows)
    build_lesson_docx(data, out)

    doc = Document(str(out))
    table = doc.tables[0]
    assert len(table.columns) == 6
    assert len(table.rows) == 6  # header + 5 data rows
    # Spot-check a few cells to confirm no shifting/truncation across columns.
    assert table.rows[1].cells[0].text == "صفأعمودأ"
    assert table.rows[1].cells[5].text == "صفأعمودو"
    assert table.rows[5].cells[5].text == "صفهـعمودو"

    errors, _warnings = validate_docx(out)
    assert errors == []


def test_header_only_table_with_zero_rows(tmp_path):
    out = tmp_path / "out.docx"
    data = _sample_with_table(["عمود وحيد"], [])
    build_lesson_docx(data, out)

    doc = Document(str(out))
    table = doc.tables[0]
    assert len(table.columns) == 1
    assert len(table.rows) == 1  # header only

    errors, _warnings = validate_docx(out)
    assert errors == []


def test_ragged_rows_are_padded_not_dropped(tmp_path):
    out = tmp_path / "out.docx"
    # Fewer cells than headers in one row, more cells than headers in another.
    data = _sample_with_table(["أ", "ب", "ج"], [["١"], ["١", "٢", "٣", "٤"]])
    build_lesson_docx(data, out)

    doc = Document(str(out))
    table = doc.tables[0]
    assert len(table.columns) == 3
    assert len(table.rows) == 3  # header + 2 data rows, neither dropped
    assert table.rows[1].cells[0].text == "١"
    assert table.rows[1].cells[1].text == ""  # padded, not misaligned
    assert table.rows[2].cells[2].text == "٣"  # extra 4th input cell silently truncated, first 3 kept


def test_table_columns_share_the_full_usable_page_width(tmp_path):
    out = tmp_path / "out.docx"
    data = _sample_with_table(["أ", "ب", "ج", "د"], [["1", "2", "3", "4"]])
    build_lesson_docx(data, out)

    doc = Document(str(out))
    table = doc.tables[0]
    expected = Inches(USABLE_WIDTH_INCHES / len(table.columns))
    for row in table.rows:
        for cell in row.cells:
            assert cell.width == expected


def test_different_column_counts_across_sections_in_one_document(tmp_path):
    out = tmp_path / "out.docx"
    data = {
        "title": "عنوان",
        "subtitle": "",
        "sections": [
            {
                "heading": "محور بجدول صغير",
                "blocks": [{"type": "paragraph", "text": "نص."}],
                "subsections": [],
                "section_summary": "",
                "table": {"title": "", "headers": ["عمود"], "rows": [["قيمة"]]},
                "concept_map": None,
            },
            {
                "heading": "محور بجدول كبير",
                "blocks": [{"type": "paragraph", "text": "نص."}],
                "subsections": [],
                "section_summary": "",
                "table": {
                    "title": "",
                    "headers": ["أ", "ب", "ج", "د", "هـ"],
                    "rows": [["1", "2", "3", "4", "5"]],
                },
                "concept_map": None,
            },
        ],
    }
    build_lesson_docx(data, out)

    doc = Document(str(out))
    assert len(doc.tables) == 2
    assert len(doc.tables[0].columns) == 1
    assert len(doc.tables[1].columns) == 5
