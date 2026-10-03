from __future__ import annotations

import zipfile

from docx import Document

from lesson_docx.builder import build_lesson_docx
from lesson_docx.validate import validate_docx

SAMPLE = {
    "title": "عنوان الدرس",
    "subtitle": "سلسلة تجريبية",
    "sections": [
        {
            "heading": "المحور الأول",
            "blocks": [
                {"type": "paragraph", "text": "نص الفقرة الأولى."},
                {"type": "bullet_list", "items": ["نقطة أولى", "نقطة ثانية"]},
                {"type": "numbered_list", "items": ["خطوة أولى", "خطوة ثانية"]},
            ],
            "subsections": [
                {
                    "heading": "عنوان فرعي",
                    "level": 2,
                    "blocks": [{"type": "paragraph", "text": "نص فرعي."}],
                    "subsections": [
                        {"heading": "عنوان من المستوى الثالث", "level": 3, "blocks": [{"type": "paragraph", "text": "نص."}], "subsections": []}
                    ],
                    "concept_map": {"label": "جذر الفرعي", "children": ["فرع أول", "فرع ثانٍ"]},
                }
            ],
            "section_summary": "خلاصة المحور الأول في جملة واحدة.",
            "table": {"title": "جدول تجريبي", "headers": ["العمود الأول", "العمود الثاني"], "rows": [["قيمة1", "قيمة2"]]},
            "concept_map": {"label": "جذر", "children": ["فرع أول", "فرع ثانٍ"]},
        },
        {
            "heading": "محور بلا خريطة ولا جدول",
            "blocks": [{"type": "paragraph", "text": "نص بسيط."}],
            "subsections": [],
            "section_summary": "",
            "table": None,
            "concept_map": None,
        },
    ],
}


def test_build_produces_valid_docx(tmp_path):
    out = tmp_path / "out.docx"
    report = build_lesson_docx(SAMPLE, out)
    assert out.exists()
    assert report.section_count == 2
    # One map on the section + one map on its level-2 subsection = two diagrams.
    assert report.smartart_success == 2
    assert report.smartart_failures == []

    errors, warnings = validate_docx(out)
    assert errors == []


def test_real_heading_styles_are_applied(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    doc = Document(str(out))
    styles_used = {p.style.name for p in doc.paragraphs if p.style is not None}

    # Title style is deliberately never used - the document title uses real
    # Heading 1 instead, so it sits atop Word's outline and can collapse the
    # whole lesson. SAMPLE nests 3 levels deep (section, level-2 subsection,
    # level-3 subsection), so Heading 4 must appear too.
    assert "Title" not in styles_used
    assert "Heading 1" in styles_used
    assert "Heading 2" in styles_used
    assert "Heading 3" in styles_used
    assert "Heading 4" in styles_used


def test_no_image_parts_are_ever_created(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    with zipfile.ZipFile(out) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert media == []


def test_section_table_and_summary_text_present(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    doc = Document(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "خلاصة المحور الأول" in all_text
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "العمود الأول"


def test_smartart_diagram_parts_present_for_section_with_concept_map(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    with zipfile.ZipFile(out) as zf:
        diagrams = [n for n in zf.namelist() if n.startswith("word/diagrams/")]
    assert any("data1" in d for d in diagrams)
