from __future__ import annotations

import zipfile

from docx import Document

from lesson_docx.builder import build_lesson_docx
from arabic_text import fix_tanween_errors, tashkeel_coverage_ratio

# Every text field carries real tashkeel, as if the AI had fully diacritized
# it - these tests confirm our own pipeline never *strips* what the model
# produced, for every field type, not just body paragraphs.
DIACRITIZED_SAMPLE = {
    "title": "فَضْلُ طَلَبِ الْعِلْمِ",
    "subtitle": "سِلْسِلَةٌ تَجْرِيبِيَّةٌ",
    "sections": [
        {
            "heading": "الْمِحْوَرُ الْأَوَّلُ",
            "blocks": [{"type": "paragraph", "text": "نَصٌّ مُشَكَّلٌ تَشْكِيلًا كَامِلًا لِلْفَقْرَةِ الْأُولَى."}],
            "subsections": [
                {
                    "heading": "عُنْوَانٌ فَرْعِيٌّ مُشَكَّلٌ",
                    "level": 2,
                    "blocks": [{"type": "paragraph", "text": "نَصٌّ فَرْعِيٌّ مُشَكَّلٌ."}],
                    "subsections": [],
                }
            ],
            "section_summary": (
                "هَذِهِ خُلَاصَةٌ مُفِيدَةٌ وَمُشَكَّلَةٌ تَشْكِيلًا كَامِلًا لِلْمِحْوَرِ الْأَوَّلِ، "
                "تَشْرَحُ لِلطَّالِبِ أَهَمَّ الْأَفْكَارِ وَالنَّتَائِجِ."
            ),
            "table": {
                "title": "جَدْوَلٌ تَوْضِيحِيٌّ مُشَكَّلٌ",
                "headers": ["الْعَمُودُ الْأَوَّلُ", "الْعَمُودُ الثَّانِي"],
                "rows": [["قِيمَةٌ أُولَى", "قِيمَةٌ ثَانِيَةٌ"]],
            },
            "concept_map": {
                "label": "الْمَفْهُومُ الرَّئِيسِيُّ الْمُشَكَّلُ",
                "children": ["فَرْعٌ أَوَّلُ مُشَكَّلٌ", "فَرْعٌ ثَانٍ مُشَكَّلٌ"],
            },
        }
    ],
}


def _build(tmp_path):
    out = tmp_path / "diacritized.docx"
    report = build_lesson_docx(DIACRITIZED_SAMPLE, out)
    assert report.smartart_success == 1, "sample must build its SmartArt cleanly for these tests to be meaningful"
    return out


def test_headings_are_diacritized_end_to_end(tmp_path):
    out = _build(tmp_path)
    doc = Document(str(out))
    heading_texts = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]

    assert DIACRITIZED_SAMPLE["title"] in heading_texts, "Heading 1 (title) must keep its diacritics verbatim"
    assert DIACRITIZED_SAMPLE["sections"][0]["heading"] in heading_texts
    assert DIACRITIZED_SAMPLE["sections"][0]["subsections"][0]["heading"] in heading_texts
    for text in heading_texts:
        assert tashkeel_coverage_ratio(text) >= 0.5


def test_table_text_is_diacritized_end_to_end(tmp_path):
    out = _build(tmp_path)
    doc = Document(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    table_spec = DIACRITIZED_SAMPLE["sections"][0]["table"]

    assert table_spec["title"] in all_text, "diacritized table title must survive verbatim"
    assert len(doc.tables) == 1
    table = doc.tables[0]
    header_cells = [c.text for c in table.rows[0].cells]
    assert header_cells == table_spec["headers"]
    data_cells = [c.text for c in table.rows[1].cells]
    assert data_cells == table_spec["rows"][0]
    for text in header_cells + data_cells:
        assert tashkeel_coverage_ratio(text) >= 0.5


def test_smartart_labels_are_diacritized_end_to_end(tmp_path):
    out = _build(tmp_path)
    with zipfile.ZipFile(out) as zf:
        data_xml = zf.read("word/diagrams/data1.xml").decode("utf-8")

    concept_map = DIACRITIZED_SAMPLE["sections"][0]["concept_map"]
    assert concept_map["label"] in data_xml, "diacritized root label must survive verbatim into the raw OOXML"
    for child_label in concept_map["children"]:
        assert child_label in data_xml, f"diacritized child label {child_label!r} must survive into data1.xml"
    assert tashkeel_coverage_ratio(concept_map["label"]) > 0.5


def test_section_summary_is_diacritized_end_to_end(tmp_path):
    out = _build(tmp_path)
    doc = Document(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    summary = DIACRITIZED_SAMPLE["sections"][0]["section_summary"]

    assert summary in all_text, "diacritized section_summary must survive verbatim"
    assert tashkeel_coverage_ratio(summary) > 0.5


# ---------------------------------------------------------------------------
# tashkeel_coverage_ratio - direct unit tests.
# ---------------------------------------------------------------------------

def test_coverage_ratio_is_high_for_fully_diacritized_text():
    assert tashkeel_coverage_ratio("الْعِلْمُ نُورٌ يَهْدِي اللَّهُ بِهِ مَنْ يَشَاءُ") > 0.5


def test_coverage_ratio_is_zero_for_undiacritized_text():
    assert tashkeel_coverage_ratio("العلم نور يهدي الله به من يشاء") == 0.0


def test_coverage_ratio_is_zero_for_empty_string():
    assert tashkeel_coverage_ratio("") == 0.0


def test_coverage_ratio_is_zero_for_non_arabic_text_and_never_raises():
    assert tashkeel_coverage_ratio("Hello world 12345") == 0.0
    assert tashkeel_coverage_ratio(None) == 0.0  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# Regression tests for the two fix_tanween_errors/_is_verb_like bugs found
# and fixed while diacritizing this round's static strings.
# ---------------------------------------------------------------------------

def test_nisba_adjective_ending_in_ya_keeps_legitimate_tanween():
    # "تَوْضِيحِيٌّ"/"تَعْلِيمِيٌّ" start with ت (a mudari'-prefix letter) but are
    # nisba adjectives, not verbs - their genuine indefinite tanween must
    # survive, not be misdetected as a verb and stripped. (Preceding word
    # deliberately does not start with ي/ت/ن/أ itself, to isolate this
    # specific fix from the separate hamza-initial-noun gap noted below.)
    assert fix_tanween_errors("جَدْوَلٌ تَوْضِيحِيٌّ") == "جَدْوَلٌ تَوْضِيحِيٌّ"
    assert fix_tanween_errors("مِثَالٌ تَعْلِيمِيٌّ") == "مِثَالٌ تَعْلِيمِيٌّ"


def test_hamza_and_nun_initial_nouns_keep_legitimate_tanween():
    """Regression test for a real bug this round's testing surfaced and fixed.

    The old `_is_verb_like` heuristic flagged any 4+-letter word starting
    with ي/ت/ن/أ (that didn't end in ي) as a verb and stripped its trailing
    tanween - but plenty of ordinary nouns start with hamza (أ) or ن and are
    4+ letters (e.g. "أُسْلُوبٌ"/"style", "نَمُوذَجٌ"/"template") without being
    verbs at all. The heuristic was removed entirely (not patched with yet
    another exception) because tanween is grammatically impossible on real
    verbs in the first place, so a word carrying it can never actually be a
    verb - see the docstring on `_is_never_tanween_eligible` in
    arabic_text.py. These nouns' legitimate tanween must now survive.
    """
    assert fix_tanween_errors("أُسْلُوبٌ") == "أُسْلُوبٌ"
    assert fix_tanween_errors("نَمُوذَجٌ") == "نَمُوذَجٌ"


def test_tanween_shadda_combination_on_final_letter_survives_either_order():
    # Tanween combined with a shadda on the word's final letter is part of
    # that letter's diacritic cluster, not "mid-word" - regardless of which
    # of the two marks the source text happens to place first.
    tanween_then_shadda = "مُسْتَحَبٌّ"  # ٌ then ّ
    shadda_then_tanween = "مُسْتَحَب" + "ّ" + "ٌ"  # ّ then ٌ (explicit, unambiguous order)
    assert fix_tanween_errors(tanween_then_shadda) == tanween_then_shadda
    assert fix_tanween_errors(shadda_then_tanween) == shadda_then_tanween
