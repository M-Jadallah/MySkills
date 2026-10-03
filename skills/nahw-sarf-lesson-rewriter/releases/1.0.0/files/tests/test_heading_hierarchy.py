from __future__ import annotations

from docx import Document

from lesson_docx.builder import build_lesson_docx
from lesson_docx.validate import validate_docx

SAMPLE = {
    "title": "عنوان الدرس الرئيسي",
    "subtitle": "",
    "sections": [
        {
            "heading": "المحور الأول",
            "blocks": [{"type": "paragraph", "text": "نص."}],
            "subsections": [
                {
                    "heading": "عنوان فرعي مستوى ٢",
                    "level": 2,
                    "blocks": [{"type": "paragraph", "text": "نص."}],
                    "subsections": [
                        {
                            "heading": "عنوان فرعي مستوى ٣",
                            "level": 3,
                            "blocks": [{"type": "paragraph", "text": "نص."}],
                            "subsections": [],
                        }
                    ],
                }
            ],
            "section_summary": "",
            "table": None,
            "concept_map": None,
        }
    ],
}


def test_title_is_the_first_paragraph_and_uses_heading_1(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    doc = Document(str(out))

    heading_paragraphs = [p for p in doc.paragraphs if p.style is not None and p.style.name.startswith("Heading")]
    assert heading_paragraphs, "expected at least one heading-styled paragraph"
    assert heading_paragraphs[0].style.name == "Heading 1"
    assert heading_paragraphs[0].text == SAMPLE["title"]


def test_only_one_heading_1_exists(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    doc = Document(str(out))

    h1_count = sum(
        1 for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"
    )
    assert h1_count == 1, "exactly one Heading 1 (the title) so the whole document nests under it"


def test_heading_sequence_is_h1_h2_h3_h4_in_document_order(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    doc = Document(str(out))

    heading_names = [
        p.style.name for p in doc.paragraphs if p.style is not None and p.style.name.startswith("Heading")
    ]
    assert heading_names == ["Heading 1", "Heading 2", "Heading 3", "Heading 4"]


def test_validate_docx_passes_on_correct_hierarchy(tmp_path):
    out = tmp_path / "out.docx"
    build_lesson_docx(SAMPLE, out)
    errors, warnings = validate_docx(out)
    assert errors == []


def test_validate_docx_errors_when_no_heading_1_present(tmp_path):
    out = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("قسم بلا عنوان رئيسي", style="Heading 2")
    doc.save(out)

    errors, _warnings = validate_docx(out)
    assert any("Heading 1" in e for e in errors)


def test_validate_docx_errors_when_heading_1_is_not_first(tmp_path):
    out = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("محور يظهر قبل العنوان الرئيسي خطأً", style="Heading 2")
    doc.add_paragraph("العنوان الرئيسي", style="Heading 1")
    doc.save(out)

    errors, _warnings = validate_docx(out)
    assert any("first heading" in e.lower() for e in errors)


def test_validate_docx_warns_on_multiple_heading_1(tmp_path):
    out = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("العنوان الأول", style="Heading 1")
    doc.add_paragraph("عنوان رئيسي آخر خطأً", style="Heading 1")
    doc.save(out)

    errors, warnings = validate_docx(out)
    assert errors == []  # a single well-formed Heading 1 at the start is fine structurally
    assert any("Heading 1" in w for w in warnings)
