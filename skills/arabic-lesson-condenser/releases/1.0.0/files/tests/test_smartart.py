from __future__ import annotations

import zipfile
from unittest import mock

from docx import Document

from lesson_docx import smartart
from lesson_docx.builder import build_lesson_docx

SAMPLE_WITH_MAP = {
    "title": "اختبار SmartArt",
    "subtitle": "",
    "sections": [
        {
            "heading": "محور بخريطة",
            "blocks": [{"type": "paragraph", "text": "نص."}],
            "subsections": [],
            "section_summary": "",
            "table": None,
            "concept_map": {"label": "جذر", "children": ["فرع"]},
        }
    ],
}


def test_successful_injection_produces_valid_smartart_structure(tmp_path):
    out = tmp_path / "out.docx"
    report = build_lesson_docx(SAMPLE_WITH_MAP, out)
    assert report.smartart_success == 1
    assert report.smartart_failures == []

    errors, warnings = smartart.validate_smartart_docx(out)
    assert errors == []

    with zipfile.ZipFile(out) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert media == [], "concept maps must never be rendered as images"


def test_injection_failure_inserts_visible_note_never_an_image(tmp_path):
    out = tmp_path / "out.docx"
    with mock.patch.object(smartart, "_read_template_part", side_effect=smartart.SmartArtError("simulated failure")):
        report = build_lesson_docx(SAMPLE_WITH_MAP, out)

    assert report.smartart_success == 0
    assert len(report.smartart_failures) == 1
    assert "simulated failure" in report.smartart_failures[0]["reason"]

    doc = Document(str(out))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "simulated failure" in all_text

    with zipfile.ZipFile(out) as zf:
        names = zf.namelist()
    assert not any(n.startswith("word/media/") for n in names), "must never fall back to an image"
    assert not any(n.startswith("word/diagrams/") for n in names), "no partial diagram parts on failure"

    # The invisible placeholder marker must have been replaced, not left dangling.
    assert not any(p.text.startswith(smartart.PLACEHOLDER_PREFIX) for p in doc.paragraphs)


def test_document_still_opens_after_failed_injection(tmp_path):
    """The atomic write must leave a valid DOCX even when inject_smartart raises mid-way."""
    out = tmp_path / "out.docx"
    with mock.patch.object(smartart, "_read_template_part", side_effect=smartart.SmartArtError("boom")):
        build_lesson_docx(SAMPLE_WITH_MAP, out)
    # Must not raise.
    Document(str(out))


# ---------------------------------------------------------------------------
# Content-adaptive structure: shape (nodes/branches/levels) and dimensions
# must scale with the tree, never a fixed/hardcoded layout.
# ---------------------------------------------------------------------------

def _wide_tree(n_children: int = 8) -> dict:
    return {"label": "الموضوع الرئيسي", "children": [f"فرع رقم {i}" for i in range(n_children)]}


def _deep_tree(depth: int = 4) -> dict:
    node = {"label": f"مستوى {depth}", "children": []}
    for level in range(depth - 1, 0, -1):
        node = {"label": f"مستوى {level}", "children": [node]}
    return node


def test_wide_tree_builds_and_validates(tmp_path):
    out = tmp_path / "out.docx"
    data = {
        "title": "اختبار خريطة واسعة",
        "subtitle": "",
        "sections": [
            {
                "heading": "محور بخريطة واسعة",
                "blocks": [{"type": "paragraph", "text": "نص."}],
                "subsections": [],
                "section_summary": "",
                "table": None,
                "concept_map": _wide_tree(8),
            }
        ],
    }
    report = build_lesson_docx(data, out)
    assert report.smartart_success == 1
    errors, _warnings = smartart.validate_smartart_docx(out)
    assert errors == []


def test_deep_tree_builds_and_validates(tmp_path):
    out = tmp_path / "out.docx"
    data = {
        "title": "اختبار خريطة عميقة",
        "subtitle": "",
        "sections": [
            {
                "heading": "محور بخريطة عميقة",
                "blocks": [{"type": "paragraph", "text": "نص."}],
                "subsections": [],
                "section_summary": "",
                "table": None,
                "concept_map": _deep_tree(4),
            }
        ],
    }
    report = build_lesson_docx(data, out)
    assert report.smartart_success == 1
    errors, _warnings = smartart.validate_smartart_docx(out)
    assert errors == []


def test_dimensions_scale_with_tree_size_but_never_exceed_page_width():
    # A tiny tree (1 leaf) sits at the base floor - width only starts growing
    # past the floor once there are enough leaves (>~20 given the current
    # leaf_w formula), so use a tree wide enough to actually exercise scaling
    # rather than one that just re-hits the same floor as the tiny tree.
    small_w, small_h = smartart._calc_dimensions(smartart._normalize_tree({"label": "x", "children": ["y"]}))
    mid_wide_w, _h = smartart._calc_dimensions(smartart._normalize_tree(_wide_tree(22)))
    deep_w, deep_h = smartart._calc_dimensions(smartart._normalize_tree(_deep_tree(4)))

    # The regression this guards against is the diagram overflowing the
    # page: the page is 8.5in wide with 0.5in margins each side = 7.5in
    # usable (6858000 EMU). The old cap (9144000 EMU / 10in) exceeded that;
    # the fixed cap must not.
    assert mid_wide_w > small_w
    assert mid_wide_w <= smartart.MAX_WIDTH_EMU
    assert smartart.MAX_WIDTH_EMU < 6858000  # strictly under the usable page width, not just under 10in

    # A deep tree should be sized taller than a tiny one, capped sanely too.
    assert deep_h > small_h
    assert deep_h <= smartart.MAX_HEIGHT_EMU


def test_very_wide_tree_is_capped_at_max_width():
    huge_w, _h = smartart._calc_dimensions(smartart._normalize_tree(_wide_tree(50)))
    assert huge_w == smartart.MAX_WIDTH_EMU


def test_concept_map_node_and_branch_count_is_content_driven():
    """The XML point count grows with the tree - it's not a fixed layout.

    Each real node contributes exactly 3 <dgm:pt> elements (the data point
    itself, a parTrans, and a sibTrans), plus one synthetic "doc" point for
    the whole diagram - this is the actual OOXML contract build_data1_xml
    implements, not an arbitrary number, so asserting it directly is a
    genuine structural check rather than a magic literal.
    """
    small_tree = smartart._normalize_tree({"label": "جذر", "children": ["فرع"]})
    wide_tree = smartart._normalize_tree(_wide_tree(8))

    small_xml = smartart.build_data1_xml(small_tree).decode("utf-8")
    wide_xml = smartart.build_data1_xml(wide_tree).decode("utf-8")

    assert small_xml.count("<dgm:pt ") == 3 * smartart._node_count(small_tree) + 1
    assert wide_xml.count("<dgm:pt ") == 3 * smartart._node_count(wide_tree) + 1
    assert wide_xml.count("<dgm:pt ") > small_xml.count("<dgm:pt ")

    # Connection count (dgm:cxn) equals node count too - one edge per node
    # linking it to its parent (the root links to the synthetic doc point).
    assert wide_xml.count("<dgm:cxn ") == smartart._node_count(wide_tree)
