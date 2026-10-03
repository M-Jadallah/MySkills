from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from . import smartart
from .styles import BLACK, BROWN, GREEN, GREEN_LIGHT, RED, add_para, add_text_runs, arabic_indic_digits, add_heading, set_paragraph_rtl, setup_document, shade_paragraph
from .tables import add_styled_table

logger = logging.getLogger(__name__)


@dataclass
class BuildReport:
    output_path: Path
    section_count: int
    smartart_success: int
    smartart_failures: list[dict[str, str]] = field(default_factory=list)


def _add_blocks(doc, blocks: list[dict[str, Any]]) -> None:
    for block in blocks:
        kind = block.get("type")
        if kind == "paragraph":
            add_para(doc, block["text"], size=22, color=BLACK, after=100)
        elif kind == "evidence":
            # Transmitted evidence (Quran verse, hadith, scholarly quote):
            # shaded light-green block with bold dark-green text so evidence
            # stands apart from the rewritten body; bracketed spans inside
            # still render brown via add_text_runs.
            p = add_para(doc, block["text"], size=22, color=GREEN, bold=True, before=40, after=100)
            shade_paragraph(p, GREEN_LIGHT)
        elif kind == "bullet_list":
            for item in block.get("items", []):
                add_para(doc, f"• {item}", size=22, color=BLACK, after=40)
        elif kind == "numbered_list":
            for i, item in enumerate(block.get("items", []), 1):
                number = arabic_indic_digits(str(i))
                add_para(doc, f"{number}. {item}", size=22, color=BLACK, after=40)


def _add_subsections(doc, subsections: list[dict[str, Any]]) -> None:
    for sub in subsections:
        add_heading(doc, sub["heading"], level=sub.get("level", 2))
        _add_blocks(doc, sub.get("blocks", []))
        _add_subsections(doc, sub.get("subsections", []))


def _add_section_summary(doc, summary_text: str) -> None:
    add_para(doc, "◆ خُلَاصَةُ القِسْمِ", size=18, color=BROWN, bold=True, before=100, after=30, keep_next=True)
    add_para(doc, summary_text, size=22, color=BLACK, bold=True, after=100)


def build_lesson_docx(data: dict[str, Any], output_path: str | Path) -> BuildReport:
    """Build the Islamic lesson DOCX from agent-authored lesson JSON.

    data shape (see references/lesson-schema.md in the skill for the full spec):
      {title, subtitle, sections: [{heading, blocks, subsections, section_summary, table, concept_map}]}
      Block types: paragraph | evidence (shaded quote block) | bullet_list | numbered_list.

    Concept maps are injected as real Word SmartArt after the initial save.
    If injection fails for a section, a visible failure note (with the
    technical reason) replaces the diagram - never an image fallback.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = setup_document()
    add_heading(doc, data["title"], level=0)
    if data.get("subtitle"):
        p = doc.add_paragraph()
        set_paragraph_rtl(p, align="center", before=0, after=160)
        add_text_runs(p, data["subtitle"], size=20, color=BROWN, bold=True)

    smartart_jobs: list[tuple[int, dict[str, Any], str]] = []
    diagram_idx = 0

    for section in data["sections"]:
        add_heading(doc, section["heading"], level=1)
        _add_blocks(doc, section.get("blocks", []))
        _add_subsections(doc, section.get("subsections", []))

        if section.get("section_summary"):
            _add_section_summary(doc, section["section_summary"])

        table = section.get("table")
        if table and table.get("headers"):
            add_styled_table(doc, table.get("title") or "جَدْوَلٌ تَوْضِيحِيٌّ", table["headers"], table.get("rows") or [])

        concept_map = section.get("concept_map")
        if concept_map:
            diagram_idx += 1
            add_para(
                doc,
                f"🧠 خَرِيطَةٌ مَفَاهِيمِيَّةٌ (SmartArt): {section['heading']}",
                size=18,
                color=RED,
                bold=True,
                before=120,
                after=40,
                keep_next=True,
            )
            placeholder_xml = smartart.make_placeholder_paragraph_xml(diagram_idx)
            placeholder_p = parse_xml(
                placeholder_xml.replace("<w:p>", f"<w:p {nsdecls('w')}>", 1)
            )
            doc.paragraphs[-1]._p.addnext(placeholder_p)
            smartart_jobs.append((diagram_idx, concept_map, section["heading"]))

    doc.save(output_path)

    smartart_failures: list[dict[str, str]] = []
    smartart_success = 0
    for idx, tree, section_heading in smartart_jobs:
        try:
            smartart.inject_smartart(output_path, tree, diagram_index=idx)
            smartart_success += 1
        except smartart.SmartArtError as exc:
            logger.warning("SmartArt injection failed for section %r (idx=%s): %s", section_heading, idx, exc)
            note = f"⚠ تَعَذَّرَ تَوْلِيدُ الخَرِيطَةِ المَفَاهِيمِيَّةِ (SmartArt) لِهَذَا المِحْوَرِ. السَّبَبُ التِّقْنِيُّ: {exc}"
            try:
                smartart.replace_placeholder_with_note(output_path, idx, note)
            except smartart.SmartArtError as note_exc:
                logger.error("Failed to insert SmartArt failure note for idx=%s: %s", idx, note_exc)
            smartart_failures.append({"section": section_heading, "reason": str(exc)})

    return BuildReport(
        output_path=output_path,
        section_count=len(data["sections"]),
        smartart_success=smartart_success,
        smartart_failures=smartart_failures,
    )
