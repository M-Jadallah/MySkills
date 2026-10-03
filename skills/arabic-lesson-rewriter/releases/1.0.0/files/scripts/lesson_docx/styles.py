from __future__ import annotations

import re

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------------------
# Unified visual identity.
#
# Primary base = islamic_lesson_bot template: Traditional Arabic, strict RTL,
# RED heading-1 / BLACK body, 0.5in margins. Table palette and structured
# table styling are carried over from islamic_study_summary_bot, since the
# lesson bot template never defined tables.
#
# Heading hierarchy uses real Word "Heading 1".."Heading 4" styles only (never
# the separate built-in "Title" style, which is not part of Word's outline
# chain and cannot have anything nest under it). The document title is
# Heading 1 - the sole top-level outline node - so collapsing it in Word's
# Navigation Pane collapses the entire lesson; major sections are Heading 2,
# and subsections are Heading 3/4, nesting correctly beneath it.
# ---------------------------------------------------------------------------

FONT = "Traditional Arabic"

BLACK = "1F1F1F"
RED = "C00000"
BROWN = "8B4513"
GREEN = "1F7A5B"
GOLD = "D9B36A"
CREAM = "FFF7E6"
GRAY = "F3F0EA"
GREEN_LIGHT = "EAF5F0"
RED_LIGHT = "FBEAEA"
WHITE = "FFFFFF"
BORDER = "D0C4B0"

BRACKET_RE = re.compile(r"(\[[^\]]+\]|\([^\)]+\)|«[^»]+»)")
ARABIC_INDIC = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def arabic_indic_digits(text: str) -> str:
    return str(text or "").translate(ARABIC_INDIC)


def get_or_add(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        # WordprocessingML property children have a prescribed order. Appending
        # bidi/spacing after existing properties produces XML that Word may
        # repair (and that strict OOXML validators reject).
        successors = {
            "w:pPr": {
                "w:keepNext": ("w:shd", "w:bidi", "w:spacing", "w:ind", "w:jc", "w:outlineLvl", "w:rPr"),
                "w:shd": ("w:bidi", "w:spacing", "w:ind", "w:jc", "w:outlineLvl", "w:rPr"),
                "w:bidi": ("w:spacing", "w:ind", "w:jc", "w:outlineLvl", "w:rPr"),
                "w:spacing": ("w:ind", "w:jc", "w:outlineLvl", "w:rPr"),
                "w:ind": ("w:jc", "w:outlineLvl", "w:rPr"),
                "w:jc": ("w:outlineLvl", "w:rPr"),
                "w:outlineLvl": ("w:rPr",),
            },
            "w:tblPr": {
                "w:bidiVisual": ("w:tblW", "w:jc", "w:tblBorders", "w:tblLook"),
            },
            "w:tcPr": {
                "w:tcBorders": ("w:shd", "w:noWrap", "w:tcMar", "w:vAlign"),
                "w:shd": ("w:noWrap", "w:tcMar", "w:vAlign"),
                "w:tcMar": ("w:vAlign",),
            },
            "w:sectPr": {
                "w:bidi": ("w:rtlGutter", "w:docGrid", "w:printerSettings"),
            },
        }
        later = successors.get("w:" + parent.tag.rsplit("}", 1)[-1], {}).get(tag, ())
        if later:
            parent.insert_element_before(child, *later)
        else:
            parent.append(child)
    return child


def _apply_rtl_font_to_rpr(rPr, *, size_half_points: str, color: str, bold: bool) -> None:
    rFonts = get_or_add(rPr, "w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)
    get_or_add(rPr, "w:sz").set(qn("w:val"), size_half_points)
    get_or_add(rPr, "w:szCs").set(qn("w:val"), size_half_points)
    get_or_add(rPr, "w:rtl").set(qn("w:val"), "1")
    get_or_add(rPr, "w:color").set(qn("w:val"), color)
    lang = get_or_add(rPr, "w:lang")
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")
    if bold:
        get_or_add(rPr, "w:b").set(qn("w:val"), "1")
        get_or_add(rPr, "w:bCs").set(qn("w:val"), "1")


def set_style_base(style, *, color: str = BLACK, bold: bool = False, size_pt: int = 22, outline: int | None = None) -> None:
    """Configure a *paragraph* style (Normal / Title / Heading N) for strict RTL Arabic.

    Setting w:outlineLvl explicitly is belt-and-suspenders: Word's built-in
    Heading N styles already carry the correct outline level, which is what
    makes the native collapse/expand feature work once this style is applied
    to a paragraph via style="Heading N".
    """
    style.font.name = FONT
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)

    rPr = style.element.get_or_add_rPr()
    _apply_rtl_font_to_rpr(rPr, size_half_points=str(size_pt * 2), color=color, bold=bold)

    pPr = style.element.get_or_add_pPr()
    get_or_add(pPr, "w:bidi")
    get_or_add(pPr, "w:jc").set(qn("w:val"), "right")
    spacing = get_or_add(pPr, "w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    if outline is not None:
        get_or_add(pPr, "w:outlineLvl").set(qn("w:val"), str(outline))


def set_paragraph_rtl(p, *, align: str = "right", before: int = 0, after: int = 120, line: int = 300, keep_next: bool = False) -> None:
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        jc_val = "center"
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        jc_val = "right"
    pPr = p._p.get_or_add_pPr()
    get_or_add(pPr, "w:bidi")
    get_or_add(pPr, "w:jc").set(qn("w:val"), jc_val)
    spacing = get_or_add(pPr, "w:spacing")
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")
    if keep_next:
        get_or_add(pPr, "w:keepNext")


def set_list_indent(p) -> None:
    """Align wrapped RTL list text under the item text, beyond its marker."""
    pPr = p._p.get_or_add_pPr()
    indent = get_or_add(pPr, "w:ind")
    indent.set(qn("w:right"), "420")
    indent.set(qn("w:hanging"), "260")


def shade_paragraph(p, fill: str) -> None:
    """Apply a background fill to a whole paragraph (evidence/quote block style)."""
    pPr = p._p.get_or_add_pPr()
    shd = get_or_add(pPr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_run_rtl(run, *, size: int = 18, color: str = BLACK, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    _apply_rtl_font_to_rpr(rPr, size_half_points=str(size * 2), color=color, bold=bold)


def add_text_runs(p, text: str, *, size: int = 18, color: str = BLACK, bold: bool = False, bracket_color: str = BROWN) -> None:
    """Split *text* on bracket-like delimiters and color bracketed spans brown."""
    text = arabic_indic_digits(text)
    for part in BRACKET_RE.split(text):
        if not part:
            continue
        is_bracket = BRACKET_RE.fullmatch(part) is not None
        run = p.add_run(part)
        set_run_rtl(run, size=size, color=bracket_color if is_bracket else color, bold=bold)


def add_para(doc_or_cell, text: str, *, size: int = 18, color: str = BLACK, bold: bool = False, align: str = "right", before: int = 0, after: int = 120, keep_next: bool = False, bracket_color: str = BROWN):
    p = doc_or_cell.add_paragraph()
    set_paragraph_rtl(p, align=align, before=before, after=after, keep_next=keep_next)
    add_text_runs(p, text, size=size, color=color, bold=bold, bracket_color=bracket_color)
    return p


#: Heading hierarchy colors, kept within the lesson-bot's own palette
#: (RED / BLACK / BROWN) rather than introducing new accent colors, so the
#: section-heading hierarchy stays faithful to the extracted template.
#: Keyed by internal level (0 = title, 1 = major section, 2/3 = subsection
#: tiers) - see add_heading() for the level -> real Word "Heading N" mapping.
_HEADING_COLORS = {0: RED, 1: RED, 2: BLACK, 3: BROWN}
_HEADING_SIZES = {0: 26, 1: 22, 2: 22, 3: 22}


def add_heading(doc: DocumentObject, text: str, *, level: int, keep_next: bool = True):
    """Add a paragraph using a REAL built-in Word heading style (Heading 1-4).

    level 0 (document title) maps to Heading 1 - the sole top-level node in
    Word's outline, so collapsing it collapses the whole lesson; level 1
    (major section) maps to Heading 2; levels 2/3 (subsection tiers) map to
    Heading 3/4. This is what makes Word's native outline collapse/expand
    work - unlike directly-formatted bold/large text, which Word cannot
    collapse, and unlike the separate built-in "Title" style, which is not
    part of the Heading 1..9 outline chain at all and cannot have subsequent
    headings nest under it.
    """
    word_level = level + 1
    style_name = f"Heading {word_level}"
    p = doc.add_paragraph(style=style_name)
    align = "center" if level == 0 else "right"
    set_paragraph_rtl(p, align=align, before=0 if level == 0 else (260 if level == 1 else 180), after=140, keep_next=keep_next)
    color = _HEADING_COLORS.get(level, BROWN)
    size = _HEADING_SIZES.get(level, 22)
    add_text_runs(p, text, size=size, color=color, bold=True)
    return p


def set_style_defaults(doc: DocumentObject) -> None:
    set_style_base(doc.styles["Normal"], color=BLACK, bold=False, size_pt=22, outline=None)
    set_style_base(doc.styles["Heading 1"], color=RED, bold=True, size_pt=26, outline=0)
    set_style_base(doc.styles["Heading 2"], color=RED, bold=True, size_pt=22, outline=1)
    set_style_base(doc.styles["Heading 3"], color=BLACK, bold=True, size_pt=22, outline=2)
    set_style_base(doc.styles["Heading 4"], color=BROWN, bold=True, size_pt=22, outline=3)


def setup_document() -> DocumentObject:
    doc = Document()
    # python-docx's bundled template omits this required attribute.
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    get_or_add(section._sectPr, "w:bidi")
    set_style_defaults(doc)
    return doc


# ---------------------------------------------------------------------------
# Table helpers (carried over near-verbatim from the study-summary template).
# ---------------------------------------------------------------------------

def shade_cell(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = get_or_add(tcPr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = get_or_add(tcPr, "w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = get_or_add(borders, f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, bottom: int = 90, left: int = 120, right: int = 120) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = get_or_add(tcPr, "w:tcMar")
    for side, value in {"top": top, "left": left, "bottom": bottom, "right": right}.items():
        node = get_or_add(tcMar, f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_rtl(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    get_or_add(tblPr, "w:bidiVisual")


def clear_cell(cell) -> None:
    cell.text = ""
    if not cell.paragraphs:
        cell.add_paragraph()


def set_cell_text(cell, text: str, *, fill: str = WHITE, color: str = BLACK, size: int = 11, bold: bool = False, align: str = "right") -> None:
    clear_cell(cell)
    shade_cell(cell, fill)
    set_cell_borders(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    set_paragraph_rtl(p, align=align, before=0, after=0, line=240)
    add_text_runs(p, text, size=size, color=color, bold=bold)
