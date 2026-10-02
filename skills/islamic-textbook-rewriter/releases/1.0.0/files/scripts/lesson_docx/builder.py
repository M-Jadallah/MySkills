"""Render textbook lessons with literal evidence, editable maps and cited closing."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches

from . import smartart
from .styles import (
    BLACK, BROWN, GREEN, GREEN_LIGHT, add_heading, add_para, add_text_runs,
    arabic_indic_digits, get_or_add, set_cell_borders, set_cell_margins,
    set_list_indent, set_paragraph_rtl, set_table_rtl, setup_document, shade_cell,
)
from .tables import add_styled_table


@dataclass
class BuildReport:
    output_path: Path
    section_count: int
    smartart_success: int = 0
    smartart_failures: list[dict[str, str]] = field(default_factory=list)


def add_link(p, url: str, label: str = "رَابِطُ الْمَصْدَرِ") -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), "Traditional Arabic")
    rpr.append(fonts)
    for tag, value in (("w:color", BROWN), ("w:sz", "32"), ("w:szCs", "32"), ("w:u", "single"), ("w:rtl", "1")):
        prop = OxmlElement(tag)
        prop.set(qn("w:val"), value)
        rpr.append(prop)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    p._p.append(hyperlink)


def add_source(doc, source: dict[str, Any]) -> None:
    parts = [source.get("author"), source["title"], source["locator"], source.get("edition")]
    p = add_para(doc, "؛ ".join(p for p in parts if p), size=16, color=BROWN, after=50)
    if source.get("url"):
        linkp = add_para(doc, "", size=16, after=60)
        add_link(linkp, source["url"])
    elif source.get("local_path"):
        # Keep paths literally; digits in filenames must not be translated.
        p = doc.add_paragraph()
        set_paragraph_rtl(p, after=60)
        add_text_runs(p, source["local_path"], size=14, color=BROWN, preserve_text=True)


def add_evidence(doc, block: dict[str, Any], sources: dict[str, dict]) -> None:
    # One cell gives a real green box enclosing quotation, citation and link.
    table = doc.add_table(rows=1, cols=1)
    set_table_rtl(table)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(7.5)
    shade_cell(cell, GREEN_LIGHT)
    set_cell_borders(cell, color=GREEN, size="8")
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)
    p = cell.paragraphs[0]
    set_paragraph_rtl(p, align="right", after=100, line=300)
    add_text_runs(p, block["text"], size=22, color=GREEN, bold=True, bracket_color=GREEN, preserve_text=True)
    p = add_para(cell, block["citation"], size=16, color=BROWN, after=40)
    source = sources[block["source_id"]]
    if source.get("url"):
        p = add_para(cell, "", size=16, after=30)
        add_link(p, source["url"])
    if block.get("note"):
        add_para(cell, block["note"], size=16, color=BROWN, after=20)
    add_para(doc, "", size=6, after=60)


def add_blocks(doc, blocks: list[dict], sources: dict[str, dict]) -> None:
    for block in blocks:
        kind = block["type"]
        if kind == "paragraph":
            add_para(doc, block["text"], size=22, color=BLACK, after=140)
        elif kind == "evidence":
            add_evidence(doc, block, sources)
        else:
            for i, item in enumerate(block["items"], 1):
                marker = "•" if kind == "bullet_list" else arabic_indic_digits(str(i)) + "."
                p = add_para(doc, marker + " " + item, size=22, after=90)
                set_list_indent(p)


def enqueue_map(doc, tree: dict, heading: str, jobs: list) -> None:
    p = doc.add_paragraph()
    set_paragraph_rtl(p, align="center", before=140, after=140)
    # A map fits as one inline drawing; no caption and no invisible spacer.
    fragment = smartart.make_placeholder_paragraph_xml(len(jobs) + 1)
    fragment = fragment.replace("<w:p>", f"<w:p {nsdecls('w')}>", 1)
    placeholder = parse_xml(fragment)
    for run in list(placeholder):
        p._p.append(run)
    jobs.append((len(jobs) + 1, tree, heading))


def render_section(doc, node: dict, depth: int, sources: dict, jobs: list) -> int:
    add_heading(doc, node["heading"], level=depth + 1)
    add_blocks(doc, node["blocks"], sources)
    if node.get("table"):
        table = node["table"]
        add_styled_table(doc, table.get("title", ""), table["headers"], table["rows"])
    total = 1
    for sub in node.get("subsections", []):
        total += render_section(doc, sub, depth + 1, sources, jobs)
    if node.get("concept_map"):
        enqueue_map(doc, node["concept_map"], node["heading"], jobs)
    return total


def build_lesson_docx(data: dict, output_path: str | Path, *, draft: bool = False) -> BuildReport:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    add_heading(doc, data["title"], level=0)
    if draft:
        add_para(doc, "مُسَوَّدَةٌ لِلْمُرَاجَعَةِ؛ لَمْ تَكْتَمِلْ بَوَّابَةُ التَّسْلِيمِ.", size=18, color=BROWN, bold=True)
    if data.get("subtitle"):
        add_para(doc, data["subtitle"], size=20, color=BROWN, align="center")
    sources = {s["id"]: s for s in data["sources"]}
    jobs: list = []
    if data.get("overview_map"):
        enqueue_map(doc, data["overview_map"], data["title"], jobs)
    count = sum(render_section(doc, section, 0, sources, jobs) for section in data["sections"])
    closing = data.get("closing")
    if closing:
        add_heading(doc, closing["heading"], level=1)
        label = "إِثْرَاءٌ تَرْبَوِيٌّ مُوَثَّقٌ"
        if closing["presentation"] == "paraphrase":
            label += " — بِالْمَعْنَى"
        add_para(doc, label, size=18, color=BROWN, bold=True, keep_next=True)
        add_blocks(doc, closing["blocks"], sources)
        add_para(doc, "الْعِبْرَةُ وَصِلَتُهَا بِالدَّرْسِ: " + closing["takeaway"], size=22, bold=True)
        add_para(doc, "الْمَصْدَرُ وَالتَّوْثِيقُ:", size=18, color=BROWN, bold=True, keep_next=True)
        for sid in closing["verification"]["source_ids"]:
            add_source(doc, sources[sid])
        add_para(doc, closing["verification"]["basis"], size=16, color=BROWN)
    elif draft:
        add_para(doc, "الْخَاتِمَةُ التَّرْبَوِيَّةُ تَحْتَاجُ إِلَى مَصْدَرٍ مُتَحَقَّقٍ مِنْهُ.", size=18, color=BROWN)
    doc.save(output_path)
    report = BuildReport(output_path, count)
    for idx, tree, heading in jobs:
        try:
            smartart.inject_smartart(output_path, tree, diagram_index=idx)
            report.smartart_success += 1
        except smartart.SmartArtError as exc:
            # Leave the temporary build for diagnosis; the CLI never publishes it.
            report.smartart_failures.append({"section": heading, "reason": str(exc)})
    return report
