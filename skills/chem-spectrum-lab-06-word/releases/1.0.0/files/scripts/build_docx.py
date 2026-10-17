#!/usr/bin/env python3
"""Chem Spectrum Lab - Step 06: build a professional Word (.docx).

Reads a step-01 lesson Markdown and produces a deep-study .docx with a
simple->complex progression, brand identity, RTL, A4, header/footer, and
colored note/warning/success boxes. Equations render to PNG via matplotlib.

Usage:
    python build_docx.py --lesson "<path>/lesson-XX.md" \
        --images "<lesson_dir>/images" --out "<lesson_dir>"
Output: <out>/lesson-XX.docx
"""
import argparse, re, pathlib, tempfile, shutil, sys

import arabic_reshaper
from bidi.algorithm import get_display

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT as WD_TBLO_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- Brand ----
PRIMARY = RGBColor(0x00, 0xD4, 0xFF)
SECONDARY = RGBColor(0x1A, 0x23, 0x7E)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEACHER = "م. معتصم جادالله (مهندس) — مدرس كيمياء"
CONTACT = "📞 0788521514  |  💬 +962 78 852 1514"
LAB = "مختبر الطيف للكيمياء"
AR_FONT = "Noto Kufi Arabic"

# box tints
NOTE_BG = "FFF9E6"; NOTE_BORDER = "D4AF37"
WARN_BG = "FFF0F0"; WARN_BORDER = "E74C3C"
OK_BG = "F0FFF4"; OK_BORDER = "2ECC71"
GOLD_BG = "FFF8E1"; GOLD_BORDER = "D4AF37"


def a(text):
    if not text:
        return ""
    try:
        return get_display(arabic_reshaper.reshape(str(text)))
    except Exception:
        return str(text)


def render_equation_png(latex, out_path):
    fig = plt.figure(figsize=(6, 1.3), dpi=200)
    fig.patch.set_alpha(0)
    ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
    ax.text(0.5, 0.5, f"${latex}$", ha="center", va="center",
            fontsize=32, color="#1A1A1A", math_fontfamily="stix")
    fig.savefig(out_path, transparent=True, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8")
        e.set(qn("w:color"), color)
        borders.append(e)
    tcPr.append(borders)


def add_boxed_para(doc, text, bg, border, icon=""):
    """A paragraph inside a single-cell shaded table = colored box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TBLO_ALIGNMENT.RIGHT
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, bg)
    set_cell_border(cell, border)
    cell.width = Inches(6.5)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(a((icon + " " if icon else "") + text))
    run.font.name = AR_FONT
    run.font.size = Pt(11)
    run.font.color.rgb = TEXT
    _rtl_run(run)
    _rtl_para(p)
    return tbl


def _rtl_run(run):
    rPr = run._element.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    rPr.append(rtl)


def _rtl_para(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)


def add_heading(doc, text, size, color, bold=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _rtl_para(p)
    r = p.add_run(a(text))
    r.font.name = AR_FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    _rtl_run(r)
    return p


def add_para(doc, text, size=12, color=TEXT, italic=False, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _rtl_para(p)
    r = p.add_run(a(text))
    r.font.name = AR_FONT
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.italic = italic
    r.font.bold = bold
    _rtl_run(r)
    return p


def add_bullet(doc, text, size=12):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _rtl_para(p)
    r = p.add_run(a("• " + text))
    r.font.name = AR_FONT
    r.font.size = Pt(size)
    r.font.color.rgb = TEXT
    _rtl_run(r)
    return p


def parse_md(path):
    text = pathlib.Path(path).read_text(encoding="utf-8")
    sections = []
    cur = {"title": "مقدمة", "level": 2, "body": []}
    for line in text.splitlines():
        m = re.match(r"^(#{2,3})\s+(.*)$", line)
        if m:
            sections.append(cur)
            cur = {"title": m.group(2).strip(), "level": len(m.group(1)), "body": []}
        else:
            cur["body"].append(line)
    sections.append(cur)
    return sections


def add_header_footer(doc, lesson_name):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    hdr = sec.header
    hp = hdr.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _rtl_para(hp)
    hr = hp.add_run(a(f"{lesson_name} — {LAB}"))
    hr.font.name = AR_FONT; hr.font.size = Pt(9); hr.font.color.rgb = SECONDARY
    _rtl_run(hr)

    ftr = sec.footer
    fp = ftr.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rtl_para(fp)
    fr = fp.add_run(a("م. معتصم جادالله  |  صفحة "))
    fr.font.name = AR_FONT; fr.font.size = Pt(9); fr.font.color.rgb = TEXT
    _rtl_run(fr)
    # page number field
    fld1 = OxmlElement("w:fldSimple"); fld1.set(qn("w:instr"), "PAGE")
    fp._p.append(fld1)


def build(doc, sections, images_dir, out, stem):
    # RTL section + A4
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.right_margin = Cm(2.5); sec.left_margin = Cm(2.5)
    # RTL
    sectPr = sec._sectPr
    bidi = OxmlElement("w:bidi")
    sectPr.append(bidi)

    # 1. Cover
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(a(f"🔬 {LAB}")); r.font.name = AR_FONT; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = PRIMARY; _rtl_run(r)
    add_para(doc, f"الدرس: {stem}", size=24, color=SECONDARY, bold=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "الصف: التوجيهي - الأردن", size=14, color=TEXT).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, TEACHER, size=14, color=SECONDARY).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, CONTACT, size=12, color=TEXT).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "الفهم قبل الحفظ", size=16, color=PRIMARY, bold=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # 2. Mind map (simple text diagram)
    add_heading(doc, "🗺️ خريطة الدرس", 18, SECONDARY)
    concepts = []
    for s in sections:
        if s["title"] not in ("معلومات الدرس", "ملاحظات المدرس", "ملاحظات إضافية"):
            concepts.append(s["title"])
    add_para(doc, "العلاقات بين مفاهيم الدرس:", size=12)
    for c in concepts[:10]:
        add_bullet(doc, c)
    add_boxed_para(doc, "خريطة الدرس توضح 'أين نحن' ضمن الموضوع الكلي للكيمياء التوجيهي.", NOTE_BG, NOTE_BORDER, "💡")
    doc.add_page_break()

    # 3. Objectives + concepts
    add_heading(doc, "🎯 الأهداف والمفاهيم الأساسية", 18, SECONDARY)
    for s in sections:
        if "الأهداف" in s["title"]:
            for line in s["body"]:
                t = line.strip()
                if t.startswith(("- ", "* ")):
                    add_bullet(doc, t[2:])
                elif re.match(r"^\d+\.", t):
                    add_bullet(doc, t)
    doc.add_page_break()

    # 4. Theory (framed paragraphs)
    add_heading(doc, "📚 المحتوى النظري", 18, SECONDARY)
    for s in sections:
        if s["title"] in ("معلومات الدرس", "ملاحظات المدرس", "ملاحظات إضافية", "الأهداف التعليمية"):
            continue
        if "المحتوى النظري" in s["title"] or s["level"] == 3:
            add_heading(doc, s["title"], 16, SECONDARY)
            for line in s["body"]:
                t = line.strip()
                if not t or t.startswith(("|", "```", "![")):
                    continue
                if t.startswith(("- ", "* ")):
                    add_bullet(doc, t[2:])
                elif re.match(r"^\d+\.", t):
                    add_para(doc, t, size=12)
                else:
                    add_para(doc, t, size=12)
                    add_boxed_para(doc, "ملاحظة مهمة: راجع هذا المفهوم جيداً فهو أساسي.", NOTE_BG, NOTE_BORDER, "⚠️")
        # embed image if referenced anywhere in section
        for line in s["body"]:
            m = re.search(r"!\[.*?\]\((.*?)\)", line)
            if m and images_dir:
                img = images_dir / pathlib.Path(m.group(1)).name
                if img.exists():
                    try:
                        doc.add_picture(str(img), width=Inches(4.5))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except Exception:
                        pass
                    break
    doc.add_page_break()

    # 5. Equations (each in its own box)
    add_heading(doc, "⚗️ المعادلات الكيميائية", 18, SECONDARY)
    eq_no = 0
    for s in sections:
        for line in s["body"]:
            for m in re.findall(r"\$\$(.+?)\$\$", line, re.S):
                eq_no += 1
                latex = m.strip()
                png = out / f"_eq_{eq_no}.png"
                render_equation_png(latex, str(png))
                add_heading(doc, f"المعادلة {eq_no}", 14, SECONDARY)
                if png.exists():
                    doc.add_picture(str(png), width=Inches(4.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_boxed_para(doc, "ملاحظة: تأكد من موازنة المعادلة.", WARN_BG, WARN_BORDER, "⚠️")
    doc.add_page_break()

    # 6. Tables (brand header)
    add_heading(doc, "📊 الجداول والمقارنات", 18, SECONDARY)
    for s in sections:
        if "جداول" in s["title"] or "جدول" in s["title"]:
            add_heading(doc, s["title"], 14, SECONDARY)
            rows = [l for l in s["body"] if l.strip().startswith("|")]
            if rows:
                data = [r.strip().strip("|").split("|") for r in rows if set(r.strip()) != {"|", "-", " "} and not set(r.strip()) <= {"|", "-", " "}]
                # filter separator rows
                data = [r for r in data if not all(set(c.strip()) <= {"-"} for c in r)]
                if data:
                    tbl = doc.add_table(rows=len(data), cols=max(len(r) for r in data))
                    tbl.style = "Table Grid"
                    for i, row in enumerate(data):
                        for j, cell in enumerate(row):
                            c = tbl.cell(i, j)
                            c.text = ""
                            para = c.paragraphs[0]; para.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _rtl_para(para)
                            run = para.add_run(a(cell.strip())); run.font.name = AR_FONT; run.font.size = Pt(11)
                            _rtl_run(run)
                            if i == 0:
                                set_cell_bg(c, "1A237E"); run.font.color.rgb = WHITE; run.font.bold = True
                            elif i % 2 == 0:
                                set_cell_bg(c, "F5F7FA")
    doc.add_page_break()

    # 7. Experiments
    add_heading(doc, "🔬 التجارب العملية", 18, SECONDARY)
    for s in sections:
        if "تجرب" in s["title"]:
            add_boxed_para(doc, s["title"], OK_BG, OK_BORDER, "🔬")
            for line in s["body"]:
                t = line.strip()
                if t.startswith(("- ", "* ")):
                    add_bullet(doc, t[2:])
                elif t:
                    add_para(doc, t, size=12)

    # 8. Solved examples
    add_heading(doc, "✏️ الأمثلة المحلولة", 18, SECONDARY)
    for s in sections:
        if "مثال" in s["title"]:
            add_boxed_para(doc, s["title"], NOTE_BG, NOTE_BORDER, "✏️")
            for line in s["body"]:
                t = line.strip()
                if t.startswith(("- ", "* ")):
                    add_bullet(doc, t[2:])
                elif re.match(r"^\d+\.", t):
                    add_para(doc, t, size=12, bold=True)
                elif t:
                    add_para(doc, t, size=12)

    # 9. Summary (golden)
    add_heading(doc, "⭐ الخلاصة (نقاط ذهبية)", 18, SECONDARY)
    add_boxed_para(doc, "أهم النقاط التي يجب تذكرها:", GOLD_BG, GOLD_BORDER, "⭐")
    for i in range(1, 6):
        add_bullet(doc, f"النقطة الذهبية {i}: [أضف أهم ما في الدرس]")

    # 10. Exercises (with writing space)
    add_heading(doc, "📝 التمارين", 18, SECONDARY)
    for s in sections:
        title_n = re.sub(r'[\u064B-\u065F\u0670]', '', s["title"])
        if not any(re.sub(r'[\u064B-\u065F\u0670]', '', k) in title_n for k in ("تمرين", "أسئلة", "مراجعة", "أتحقق")):
            continue
        for line in s["body"]:
            t = line.strip()
            if t.startswith(("###", "##")):
                add_heading(doc, t.lstrip("#").strip(), 14, SECONDARY)
                continue
            if not t:
                continue
            if t.startswith("**"):
                t2 = t.replace("*", "").strip()
                add_para(doc, t2, size=12, bold=True)
                sp = doc.add_paragraph(); sp.add_run().add_break(); sp.add_run().add_break()
            elif t.startswith(("- ", "* ")):
                t2 = t[2:].replace("*", "").strip()
                add_para(doc, t2, size=12, bold=True)
                sp = doc.add_paragraph(); sp.add_run().add_break(); sp.add_run().add_break()
            elif re.match(r"^\d+\.", t):
                add_para(doc, t, size=12, bold=True)
                sp = doc.add_paragraph(); sp.add_run().add_break(); sp.add_run().add_break()
            else:
                add_para(doc, t, size=12)
    add_boxed_para(doc, "💡 تلميح: راجع المعادلات والجداول أعلاه قبل الحل.", NOTE_BG, NOTE_BORDER, "💡")

    # 11. Glossary
    add_heading(doc, "📖 قاموس المصطلحات", 18, SECONDARY)
    tbl = doc.add_table(rows=1, cols=2); tbl.style = "Table Grid"
    h1 = tbl.cell(0, 0); h2 = tbl.cell(0, 1)
    for c, txt in ((h1, "المصطلح"), (h2, "التعريف")):
        c.text = ""; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _rtl_para(p)
        r = p.add_run(a(txt)); r.font.name = AR_FONT; r.font.bold = True; r.font.color.rgb = WHITE; _rtl_run(r)
        set_cell_bg(c, "1A237E")
    for s in sections:
        if "مصطلح" in s["title"]:
            for line in s["body"]:
                t = line.strip()
                if t.startswith("|") and "---" not in t:
                    parts = [x.strip() for x in t.strip("|").split("|")]
                    if len(parts) >= 2 and parts[0] and parts[0] != "المصطلح":
                        row = tbl.add_row().cells
                        for ci, val in enumerate(parts[:2]):
                            row[ci].text = ""; p = row[ci].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; _rtl_para(p)
                            r = p.add_run(a(val)); r.font.name = AR_FONT; r.font.size = Pt(11); _rtl_run(r)

    # 12. Contacts
    add_heading(doc, "🔗 روابط مفيدة وتواصل", 18, SECONDARY)
    add_para(doc, f"موقع المختبر: {LAB}", size=12, color=PRIMARY)
    add_para(doc, CONTACT, size=12)

    # header/footer
    add_header_footer(doc, stem)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lesson", required=True)
    ap.add_argument("--images", default="")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    out = pathlib.Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"\.md$", "", pathlib.Path(args.lesson).stem)
    docx_path = out / f"{stem}.docx"
    images_dir = pathlib.Path(args.images) if args.images else None

    sections = parse_md(args.lesson)
    doc = Document()
    build(doc, sections, images_dir, out, stem)
    doc.save(str(docx_path))

    # cleanup temp eq pngs
    for p in out.glob("_eq_*.png"):
        try: p.unlink()
        except Exception: pass
    print(f"DOCX : {docx_path}")


if __name__ == "__main__":
    main()
